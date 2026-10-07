"""The Full Employee Listing's supporting sheets.

Summary - Basis of Cover, Declaration, the System Category grid (how roster
attributes map onto each product's category), Headcount Summary (live
COUNTIFS/SUMIFS over the listing) and Setup & Data Gaps.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date
from typing import Any

from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from app.models import BrokerFirm
from app.services.full_el.context import ElContext
from app.services.full_el.listing_sheet import FIRST_ROW, employee_covers
from app.services.insurer_reports import safe_cell
from app.services.roster_attributes import roster_date

LISTING = "Employee Listing"
_BOLD = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")


def _products_in_block(ctx: ElContext, index: int) -> list[Any]:
    return [ctx.products[c] for c in ctx.block_products.get(index, []) if c in ctx.products]


def _block_labels(ctx: ElContext, index: int) -> list[tuple[str, Any]]:
    """(listing label, category info) for the block's categories, in order."""
    codes = set(ctx.block_products.get(index, []))
    seen: dict[str, Any] = {}
    for info in ctx.categories.values():
        if info.product.code in codes:
            seen.setdefault(info.label, info)
    return sorted(seen.items(), key=lambda item: item[0])


def write_summary(ws: Worksheet, ctx: ElContext) -> None:
    ws["A1"], ws["A1"].font = f"{ctx.client.legal_name or ctx.client.name} — Basis of Cover", _BOLD
    ws["A2"] = f"Benefit year {ctx.year.start_date:%d/%m/%Y} to {ctx.year.end_date:%d/%m/%Y}"
    row = 4
    for col, title in zip("ABCDEFGH", ("No.", "Product", "Insurer", "Policy No.", "Period",
                                       "Administration", "Premium currency", "Categories"),
                          strict=True):
        ws[f"{col}{row}"], ws[f"{col}{row}"].font = title, _BOLD
    row += 1
    number = 0
    for index, block in enumerate(ctx.layout.blocks):
        if block.kind != "product":
            continue
        number += 1
        products = _products_in_block(ctx, index)
        ws[f"A{row}"], ws[f"B{row}"] = number, block.banner
        ws[f"B{row}"].font = _BOLD
        for info in products:
            start, end = info.period
            ws[f"C{row}"] = safe_cell(info.insurer) or None
            ws[f"D{row}"] = safe_cell(info.policy_no) or None
            ws[f"E{row}"] = (f"{start:%d/%m/%Y} to {end:%d/%m/%Y}" if start and end else None)
            # The reviewed interpretation sits under the slip's administration wording.
            admin = "\n".join(a for a in (info.admin_basis, info.rules.admin_resolution) if a)
            ws[f"F{row}"] = safe_cell(admin) or None
            ws[f"F{row}"].alignment = _WRAP
            ws[f"G{row}"] = safe_cell(info.rules.currency) or None
            codes = {info.product.code}
            labels = [lbl for lbl, cat in _block_labels(ctx, index) if cat.product.code in codes]
            ws[f"H{row}"] = "\n".join(labels) or "NIL"
            ws[f"H{row}"].alignment = _WRAP
            row += 1
        row += 1
    for col, width in zip("ABCDEFGH", (6, 40, 22, 14, 26, 30, 12, 70), strict=True):
        ws.column_dimensions[col].width = width


def write_declaration(ws: Worksheet, ctx: ElContext, firm: BrokerFirm | None) -> None:
    broker = firm.name if firm is not None else "the broker"
    lines = (
        "By completing and submitting this file, you declare that",
        f"i) All information submitted or to be submitted is accurate and complete, and you "
        f"will inform {broker} if there are any changes to the data provided.",
        f"ii) You have received the consent from 3rd parties to provide their data to {broker} "
        "in relation to the relevant purposes.",
    )
    for i, line in enumerate(lines, start=1):
        ws[f"A{i}"], ws[f"A{i}"].alignment = line, _WRAP
    ws.column_dimensions["A"].width = 120


def write_system_category(ws: Worksheet, ctx: ElContext) -> bool:
    """Grade and work pass -> each block's category, as the listing applies it."""
    product_blocks = [(i, b) for i, b in enumerate(ctx.layout.blocks) if b.kind == "product"]
    grid: dict[tuple[str, str], list[Counter[str]]] = defaultdict(
        lambda: [Counter() for _ in product_blocks]
    )
    deps_with_listing = {key[2:] for key, _ in ctx.listed if key.startswith("D:")}
    for emp in ctx.employees:
        attrs = emp.attribute_values or {}
        grade, work_pass = str(attrs.get("grade") or ""), str(attrs.get("pass") or "")
        if not grade and not work_pass:
            continue
        covers = employee_covers(ctx, emp, deps_with_listing)
        counters = grid[(grade, work_pass)]
        for slot, (index, _block) in enumerate(product_blocks):
            cover = covers.get(index)
            counters[slot][cover.info.label if cover else "Not covered"] += 1
    if not grid:
        return False
    headers = ["Grade", "Work pass", "Employees", *(b.banner for _, b in product_blocks)]
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(1, col, title)
        cell.font, cell.alignment = _BOLD, _WRAP
    for row, ((grade, work_pass), counters) in enumerate(sorted(grid.items()), start=2):
        ws.cell(row, 1, safe_cell(grade) or None)
        ws.cell(row, 2, safe_cell(work_pass) or None)
        ws.cell(row, 3, sum(counters[0].values()) if counters else 0)
        for slot, counter in enumerate(counters):
            parts = [f"{label} ({n})" if len(counter) > 1 else label
                     for label, n in counter.most_common()]
            ws.cell(row, 4 + slot, safe_cell("; ".join(parts))).alignment = _WRAP
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 12 if col <= 3 else 38
    return True


def write_headcount(ws: Worksheet, ctx: ElContext, last_row: int) -> None:
    def rng(letter: str) -> str:
        return f"'{LISTING}'!${letter}${FIRST_ROW}:${letter}${max(last_row, FIRST_ROW)}"

    layout = ctx.layout
    dep_name = layout.dependant.column("name") if layout.dependant else None
    row = 1
    for index, block in enumerate(layout.blocks):
        if block.kind != "product":
            continue
        category = block.column("category")
        if category is None:
            continue
        ws.cell(row, 1, block.banner).font = _BOLD
        row += 1
        titles = ["Category", "Participation", "Basis", "No. of employees",
                  "No. of dependants", "Sum insured", "Premium", "Premium with GST"]
        for col, title in enumerate(titles, start=1):
            ws.cell(row, col, title).font = _BOLD
        row += 1
        first = row
        for label, info in _block_labels(ctx, index):
            crit = f'"{label.replace(chr(34), chr(34) * 2)}"'
            ws.cell(row, 1, safe_cell(label))
            ws.cell(row, 2, info.category.participation_model)
            pa = info.category.plan_assignments or {}
            ws.cell(row, 3, safe_cell(str(pa.get("basis") or "")) or None)
            if dep_name is not None:
                ws.cell(row, 4, f'=COUNTIFS({rng(category.letter)},{crit},'
                                f'{rng(dep_name.letter)},"")')
                ws.cell(row, 5, f'=COUNTIFS({rng(category.letter)},{crit},'
                                f'{rng(dep_name.letter)},"<>")')
            else:
                ws.cell(row, 4, f"=COUNTIFS({rng(category.letter)},{crit})")
            for col, role in ((6, "eligible_si"), (7, "premium"), (8, "premium_gst")):
                target = block.column(role)
                if target is not None:
                    ws.cell(row, col, f"=SUMIFS({rng(target.letter)},"
                                      f"{rng(category.letter)},{crit})")
                    ws.cell(row, col).number_format = "#,##0.00"
            row += 1
        ws.cell(row, 1, "Total").font = _BOLD
        for col in range(4, 9):
            letter = ws.cell(row, col).column_letter
            total = ws.cell(row, col, f"=SUM({letter}{first}:{letter}{row - 1})")
            total.font = _BOLD
            total.number_format = "#,##0" if col < 6 else "#,##0.00"
        row += 2
    for letter, width in zip("ABCDEFGH", (60, 14, 28, 16, 16, 18, 16, 18), strict=True):
        ws.column_dimensions[letter].width = width


def setup_gaps(ctx: ElContext, stats: dict[str, Any]) -> list[tuple[str, str, str]]:
    gaps = list(ctx.gaps)
    for code, info in sorted(ctx.products.items()):
        if not info.confirmed:
            gaps.append((code, "Setup not confirmed",
                         "Rates and terms come from the uploaded slip; confirm the setup."))
        if info.gst_factor is None:
            gaps.append((code, "GST unknown",
                         "Premium with GST is left blank; set GST on the product terms."))
        if info.rules.age_reference == "Member effective date" and (missing := sum(
            not isinstance(roster_date((e.attribute_values or {}).get("effective_date")), date)
            for e in ctx.employees
        )):
            gaps.append((code, f"{missing} employees without an effective date",
                         "Ages for this product run to the member effective date; "
                         "theirs are left blank."))
    for cat in ctx.categories.values():
        if cat.rate_basis == "earnings_based":
            continue
        if not (cat.rate or cat.tiers):
            gaps.append((cat.product.code, f"No rate: {cat.category.display_name}",
                         "Premium is left blank; enter the rate in the product setup."))
    unmatched = sum(1 for e in ctx.employees if not e.matched_categories)
    if unmatched:
        gaps.append(("All", f"{unmatched} employees without cover",
                     "Listed with no cover, or not matched to any category."))
    if stats.get("missing_salary"):
        gaps.append(("All", f"{stats['missing_salary']} employees without salary",
                     "Salary-multiple sums insured cannot be calculated for them."))
    if stats.get("missing_dob"):
        gaps.append(("All", f"{stats['missing_dob']} employees without date of birth",
                     "Ages, age limits and underwriting checks cannot be applied."))
    if stats.get("uw"):
        gaps.append(("All", f"{stats['uw']} employees above the non-evidence limit",
                     "New cover or an increase above the NEL amount or age needs underwriting."))
    return gaps


def write_gaps(ws: Worksheet, gaps: list[tuple[str, str, str]]) -> None:
    for col, title in enumerate(("Product", "Finding", "Effect / action"), start=1):
        ws.cell(1, col, title).font = _BOLD
    for row, (product, finding, action) in enumerate(gaps, start=2):
        ws.cell(row, 1, safe_cell(product))
        ws.cell(row, 2, safe_cell(finding))
        ws.cell(row, 3, safe_cell(action)).alignment = _WRAP
    for letter, width in zip("ABC", (16, 50, 80), strict=True):
        ws.column_dimensions[letter].width = width
