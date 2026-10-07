"""The "Employee Listing" sheet: one row per employee, then their dependants."""
from __future__ import annotations

from datetime import date
from typing import Any

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import Employee
from app.services.el_workbook import ElBlock
from app.services.full_el.cells import (
    DATE_ROLES,
    Cover,
    Refs,
    dependant_value,
    employee_value,
    needs_underwriting,
    number_text,
    product_age_basis,
    product_value,
)
from app.services.full_el.context import ElContext

DATE_FORMAT = "dd/mmm/yy"
MONEY_FORMAT = "#,##0.00"
FIRST_ROW = 4  # rows 1-3: reference cells, banners, headers
_MONEY_ROLES = {"salary", "present_si", "eligible_si", "pending_si", "last_accepted_si",
                "premium", "premium_gst"}


def _anb(dob: Any, reference: date) -> int | None:
    try:
        born = date.fromisoformat(str(dob)[:10])
    except ValueError:
        return None
    before_birthday = (reference.month, reference.day) < (born.month, born.day)
    return reference.year - born.year - before_birthday + 1


def _family_group(relationships: list[str]) -> str:
    spouse = any(r.lower().startswith("s") for r in relationships)
    child = any(r.lower().startswith("c") for r in relationships)
    return {(False, False): "EO", (True, False): "ES", (False, True): "EC"}.get(
        (spouse, child), "EF"
    )


def employee_covers(ctx: ElContext, emp: Employee, deps_with_listing: set[str]) -> dict[int, Cover]:
    matched = {
        m.get("product_code"): m.get("category_id")
        for m in emp.matched_categories or [] if isinstance(m, dict)
    }
    out: dict[int, Cover] = {}
    for block, codes in ctx.block_products.items():
        info = next(
            (ctx.categories[cid] for code in codes
             if (cid := matched.get(code)) and cid in ctx.categories),
            None,
        )
        if info is None:
            continue
        product_id = info.product.id
        covered = []
        for dep in ctx.dependants.get(emp.id, []):
            if (f"D:{dep.id}", product_id) in ctx.listed or (
                dep.id not in deps_with_listing and info.covers_dependants
                and dep.status == "active"
            ):
                covered.append(dep)
        relationships = [str((d.attribute_values or {}).get("relationship") or "") for d in covered]
        out[block] = Cover(
            info=info,
            assignment=ctx.listed.get((f"E:{emp.id}", product_id)),
            family_group=_family_group(relationships),
            covered_dependants=covered,
        )
    return out


def ordered_employees(ctx: ElContext) -> list[Employee]:
    """The listing's own row order where known, then staff ID."""
    rows: dict[str, int] = {}
    for (key, _product), a in ctx.listed.items():
        if key.startswith("E:") and a.source_row:
            rows[key[2:]] = min(rows.get(key[2:], a.source_row), a.source_row)
    return sorted(ctx.employees, key=lambda e: (rows.get(e.id, 10**9), e.staff_id or ""))


def _rate_cells(ctx: ElContext) -> dict[int, tuple[str, float]]:
    """Blocks rated per $1,000 SI at one rate get that rate in row 1."""
    out: dict[int, tuple[str, float]] = {}
    for index, block in enumerate(ctx.layout.blocks):
        premium = block.column("premium")
        codes = set(ctx.block_products.get(index, []))
        rates = {
            info.rate for info in ctx.categories.values()
            if info.product.code in codes and info.rate_basis == "per_1000_si" and info.rate
        }
        if premium is not None and len(rates) == 1:
            out[index] = (premium.letter, rates.pop())
    return out


def _header_text(ctx: ElContext, index: int, header: str, role: str) -> str:
    """Age headings name their basis; a product block only when it differs."""
    if role != "age" or header.strip() != "Age":
        return header
    block = ctx.layout.blocks[index]
    if block.kind in {"employee", "dependant"}:
        return f"Age ({ctx.age_basis})"
    bases = {product_age_basis(ctx.products[code], ctx)
             for code in ctx.block_products.get(index, []) if code in ctx.products}
    if len(bases) == 1 and bases != {ctx.age_basis}:
        return f"Age ({bases.pop()})"
    return header


def write_listing(ws: Worksheet, ctx: ElContext, *, masked: bool) -> dict[str, Any]:
    layout = ctx.layout
    emp_block, dep_block = layout.employee, layout.dependant

    def letter(block: ElBlock | None, role: str) -> str | None:
        column = block.column(role) if block else None
        return column.letter if column else None

    dob = letter(emp_block, "date_of_birth")
    refs = Refs(
        dob=dob, dep_dob=letter(dep_block, "date_of_birth"),
        age=letter(emp_block, "age"), dep_age=letter(dep_block, "age"),
        salary=letter(emp_block, "salary"),
        reference=f"${dob}$1" if dob else None,
    )
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    if dob:
        ws[f"{dob}1"] = ctx.reference_date
        ws[f"{dob}1"].number_format = DATE_FORMAT
    rate_cells = _rate_cells(ctx)
    for letter_, rate in rate_cells.values():
        ws[f"{letter_}1"] = f"={number_text(rate)}/1000"

    for index, block in enumerate(layout.blocks):
        anchor = block.column("category") if block.kind == "product" else block.column("name")
        anchor = anchor or (block.columns[0] if block.columns else None)
        if anchor is not None:
            ws[f"{anchor.letter}2"] = block.banner
            ws[f"{anchor.letter}2"].font = bold
        for column in block.columns:
            cell = ws[f"{column.letter}3"]
            cell.value = _header_text(ctx, index, column.header, column.role)
            cell.font, cell.alignment = bold, wrap
    for column in layout.trailing:
        cell = ws[f"{column.letter}3"]
        cell.value, cell.font, cell.alignment = column.header, bold, wrap

    deps_with_listing = {key[2:] for key, _ in ctx.listed if key.startswith("D:")}
    stats = {"employees": 0, "dependants": 0, "uw": 0, "missing_salary": 0, "missing_dob": 0}
    row = FIRST_ROW
    for emp in ordered_employees(ctx):
        covers = employee_covers(ctx, emp, deps_with_listing)
        attrs = emp.attribute_values or {}
        salary = _number(attrs.get("salary"))
        anb = _anb(attrs.get("date_of_birth"), ctx.reference_date)
        stats["employees"] += 1
        if anb is None:
            stats["missing_dob"] += 1
        if salary is None and any(c.info.salary_multiple for c in covers.values()):
            stats["missing_salary"] += 1
        uw = {b: needs_underwriting(c, ctx, salary, anb) for b, c in covers.items()}
        stats["uw"] += any(uw.values())
        for column in [*(emp_block.columns if emp_block else ()), *layout.trailing]:
            _put(ws, column.letter, row, column.role,
                 employee_value(column.role, column.letter, emp, ctx, row, refs, masked))
        _write_covers(ws, ctx, covers, row, refs, rate_cells, uw, None, attrs)
        row += 1
        for dep in ctx.dependants.get(emp.id, []):
            if dep_block is None:
                break
            # A dependant row repeats the staff ID and name to link to its employee.
            for role in ("staff_id", "name"):
                link = emp_block.column(role) if emp_block else None
                if link is not None:
                    _put(ws, link.letter, row, role,
                         employee_value(role, link.letter, emp, ctx, row, refs, masked))
            for column in dep_block.columns:
                _put(ws, column.letter, row, column.role,
                     dependant_value(column.role, dep, ctx, row, refs, masked))
            dep_covers = {
                b: Cover(c.info, ctx.listed.get((f"D:{dep.id}", c.info.product.id)), None, [])
                for b, c in covers.items() if dep in c.covered_dependants
            }
            _write_covers(ws, ctx, dep_covers, row, refs, rate_cells, uw, dep,
                          dep.attribute_values or {})
            stats["dependants"] += 1
            row += 1
    stats["last_row"] = row - 1

    for column in [c for b in layout.blocks for c in b.columns] + list(layout.trailing):
        ws.column_dimensions[column.letter].width = 14 if column.role not in {
            "name", "category", "remarks", "designation"} else 30
    name_column = emp_block.column("name") if emp_block else None
    freeze_col = get_column_letter(name_column.index + 2) if name_column else "A"
    ws.freeze_panes = f"{freeze_col}{FIRST_ROW}"
    return stats


def _write_covers(
    ws: Worksheet, ctx: ElContext, covers: dict[int, Cover], row: int, refs: Refs,
    rate_cells: dict[int, tuple[str, float]], uw: dict[int, bool], dependant: Any,
    attrs: dict[str, Any],
) -> None:
    for index, cover in covers.items():
        block = ctx.layout.blocks[index]
        letters = {c.role: c.letter for c in block.columns}
        rate_cell = (f"${rate_cells[index][0]}$1"
                     if index in rate_cells and cover.info.rate == rate_cells[index][1] else None)
        for column in block.columns:
            value = product_value(
                column.role, block, letters, cover, ctx, row, refs,
                dependant=dependant, attrs=attrs, rate_cell=rate_cell,
                uw=uw.get(index, False),
            )
            _put(ws, column.letter, row, column.role, value)


def _put(ws: Worksheet, letter: str, row: int, role: str, value: Any) -> None:
    if value is None:
        return
    cell = ws[f"{letter}{row}"]
    cell.value = value
    if role in DATE_ROLES and isinstance(value, date):
        cell.number_format = DATE_FORMAT
    elif role in _MONEY_ROLES:
        cell.number_format = MONEY_FORMAT


def _number(value: Any) -> float | None:
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None

