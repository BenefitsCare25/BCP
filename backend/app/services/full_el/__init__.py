"""Full Employee Listing (EL) in the company's own template layout.

Sheets: Summary - Basis of Cover, Employee Listing (employee rows followed by
their dependants, one column block per product, live formulas), Declaration,
the System Category grid, Headcount Summary and Setup & Data Gaps.
"""
from __future__ import annotations

from openpyxl import Workbook
from sqlalchemy.orm import Session

from app.models import BrokerFirm, PolicyYear
from app.services.brand import DEFAULT_BRAND, resolve_brand
from app.services.full_el.context import load_context
from app.services.full_el.listing_sheet import write_listing
from app.services.full_el.sheets import (
    LISTING,
    setup_gaps,
    write_declaration,
    write_gaps,
    write_headcount,
    write_summary,
    write_system_category,
)
from app.services.report_workbooks import SYSTEM_CATEGORY_SHEET, branded_title


def build_full_el(
    db: Session, year: PolicyYear, *, masked: bool = True, employee_status: str = "all"
) -> Workbook:
    ctx = load_context(db, year, employee_status=employee_status)
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary - Basis of Cover"
    listing = wb.create_sheet(LISTING)
    stats = write_listing(listing, ctx, masked=masked)
    write_summary(summary, ctx)
    firm = db.get(BrokerFirm, ctx.client.broker_firm_id) if ctx.client.broker_firm_id else None
    write_declaration(wb.create_sheet("Declaration - Please Read"), ctx, firm)
    brand = resolve_brand(db, firm.id) if firm is not None else DEFAULT_BRAND
    system = wb.create_sheet(branded_title(SYSTEM_CATEGORY_SHEET, brand.short_name))
    if not write_system_category(system, ctx):
        wb.remove(system)
    write_headcount(wb.create_sheet("Headcount Summary"), ctx, stats["last_row"])
    write_gaps(wb.create_sheet("Setup & Data Gaps"), setup_gaps(ctx, stats))
    return wb


__all__ = ["build_full_el"]
