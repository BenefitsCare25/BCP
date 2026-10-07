"""Locate an Employee Listing's banner/header rows and the role of each column."""
from __future__ import annotations

import re
from collections.abc import Sequence

from openpyxl.utils import get_column_letter

from app.services import product_registry
from app.services.el_workbook.models import BlockKind, ElBlock, ElColumn, ElLayout
from app.services.el_workbook.values import as_date
from app.services.excel_reader import Cell

# (role, pattern) checked in order; the first match wins. Patterns run on the
# lower-cased header with whitespace collapsed.
_EMPLOYEE_ROLES: tuple[tuple[str, str], ...] = (
    ("movement", r"\baddition\b|\bdeletion\b|^a\s*-\s*addition"),
    ("staff_id", r"(employee|staff)\s*(id|no|number|code)"),
    ("national_id", r"\bnric\b|\bfin\b|passport|national\s*id"),
    ("date_of_birth", r"date\s*of\s*birth|\bdob\b"),
    ("date_of_hire", r"date\s*of\s*(hire|join)|join(ed|ing)?\s*date|employment\s*date"),
    ("name", r"^name\b|employee\s*name|full\s*name"),
    ("gender", r"^sex\b|gender"),
    ("department", r"^dept\b|department"),
    ("designation", r"designation|job\s*title|position"),
    ("grade", r"grade"),
    ("citizenship", r"citizenship|residen(cy|tial)\s*status|pass\s*type|work\s*pass"),
    ("nationality", r"nationality"),
    ("age", r"^age\b"),
    ("location", r"depot|work\s*location|^location|^site\b"),
    ("salary", r"salary|basic\s*pay|monthly\s*wage"),
    ("bank_code", r"bank\s*code"),
    ("bank_branch", r"branch\s*code"),
    ("bank_account", r"bank\s*account|account\s*no"),
    ("email", r"e-?mail"),
    ("mobile", r"mobile|phone|contact\s*no"),
    ("entity", r"^entity\b|^company\b|insured\s*entity"),
    ("country_of_work", r"country\s*of\s*work"),
    ("marital_status", r"marital\s*status"),
    ("cost_centre", r"cost\s*cent(re|er)"),
    ("currency", r"^currency\b"),
    ("category", r"category"),
)

_DEPENDANT_ROLES: tuple[tuple[str, str], ...] = (
    ("relationship", r"relationship"),
    ("national_id", r"\bnric\b|\bfin\b|passport|national\s*id"),
    ("date_of_birth", r"date\s*of\s*birth|\bdob\b"),
    ("name", r"name"),
    ("gender", r"^sex\b|gender"),
    ("nationality", r"nationality"),
    ("residence", r"country\s*of\s*residence|residence"),
    ("age", r"^age\b"),
)

_PRODUCT_ROLES: tuple[tuple[str, str], ...] = (
    ("admin_type", r"type\s*of\s*administration|headcount\s*/\s*named"),
    ("pending_si", r"pending\s*sum\s*insured"),
    ("last_accepted_si", r"last\s*accepted\s*sum\s*insured"),
    ("present_si", r"present\s*sum\s*insured|current\s*sum\s*insured"),
    ("eligible_si", r"eligible\s*sum\s*insured|sum\s*insured"),
    ("mu_status", r"mu\s*\)?\s*status|\(mu\s*status|underwriting\s*status"),
    ("mu_decision", r"mu\s*decision|underwriting\s*decision"),
    ("mu_letter_insurer", r"letter\s*sent\s*date\s*to\s*insurer"),
    ("mu_letter_member", r"letter\s*sent\s*date"),
    ("acceptance_date", r"acceptance\s*date"),
    ("loading_rate", r"h\s*&\s*/?\s*or\s*r|loading"),
    ("premium_gst", r"premium.*(with|incl\.?|including)\s*gst"),
    ("premium", r"premium"),
    ("plan_type", r"plan\s*type|\(plan"),
    ("family_group", r"family\s*group|\beo\s*/\s*es"),
    ("category", r"category"),
    ("age", r"^age\b"),
)

_TRAILING_ROLES: tuple[tuple[str, str], ...] = (
    ("last_day_of_service", r"last\s*day\s*of\s*service|date\s*of\s*(exit|leaving|termination)"),
    ("remarks", r"^remarks?\b"),
    ("internal_note", r"internal\s*use|updated\s*date\s*of\s*changes"),
    ("retrenchment_start", r"retrenchment\s*commencement"),
    ("retrenchment_end", r"retrenchment\s*end"),
    ("reemployment_start", r"re-?employment\s*commencement"),
    ("reemployment_end", r"re-?employment\s*end"),
)

_HEADER_SIGNALS = (
    r"name", r"nric|fin", r"date\s*of\s*birth", r"category", r"(employee|staff)\s*id",
)


def norm_header(value: Cell) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _role(header: str, table: Sequence[tuple[str, str]]) -> str:
    low = header.lower()
    return next((role for role, pattern in table if re.search(pattern, low)), "other")


def _find_header_row(rows: list[list[Cell]]) -> int:
    best, best_score = -1, 0
    for i, row in enumerate(rows[:15]):
        text = " | ".join(norm_header(c).lower() for c in row if c not in (None, ""))
        score = sum(bool(re.search(p, text)) for p in _HEADER_SIGNALS)
        if score > best_score:
            best, best_score = i, score
    return best if best_score >= 3 else -1


def _banner_kind(text: str) -> BlockKind:
    low = text.lower()
    if re.search(r"depend[ae]n[td]|family\s*member", low):
        return "dependant"
    if re.fullmatch(r"\s*(employees?|members?|staff)\s*(details?)?\s*", low):
        return "employee"
    return "product"


def _code_hint(banner: str, headers: list[str]) -> str | None:
    known = product_registry.known_codes()
    for token in re.findall(r"\(([A-Za-z]{2,6})\)", banner):
        if token.upper() in known:
            return token.upper()
    for header in headers:
        match = re.match(r"\s*([A-Za-z]{2,6})\b", header)
        if match and match.group(1).upper() in known and "premium" in header.lower():
            return match.group(1).upper()
    return None


def _leading_product_start(header: list[Cell], banner_col: int) -> int:
    """First column of a product block whose banner sits over its Category.

    Listings often print a product's "Type of Administration" and "Age"
    columns just left of the banner. Those belong to the product, so the block
    starts at the administration column. An age column without one stays with
    the block before it (it is usually the dependant's age).
    """
    taken: list[int] = []
    for c in range(banner_col - 1, max(banner_col - 3, -1), -1):
        role = _role(norm_header(header[c]) if c < len(header) else "", _PRODUCT_ROLES)
        if role not in {"admin_type", "age"}:
            break
        taken.append(c)
        if role == "admin_type":
            return c
    return banner_col


def _columns(
    header: list[Cell], start: int, stop: int, table: Sequence[tuple[str, str]]
) -> list[ElColumn]:
    out: list[ElColumn] = []
    for c in range(start, stop):
        text = norm_header(header[c]) if c < len(header) else ""
        if not text:
            continue
        trailing = _role(text, _TRAILING_ROLES)
        if trailing != "other":
            continue
        out.append(ElColumn(c, get_column_letter(c + 1), text, _role(text, table)))
    return out


def detect_layout(sheet_name: str, rows: list[list[Cell]]) -> ElLayout | None:
    """The sheet's EL layout, or None when it isn't an employee listing."""
    header_row = _find_header_row(rows)
    if header_row < 0:
        return None
    header = rows[header_row]
    width = len(header)
    banner_row = header_row - 1
    banner = rows[banner_row] if banner_row >= 0 else []
    starts = [
        (c, norm_header(v)) for c, v in enumerate(banner) if norm_header(v) and c < width
    ]
    if len(starts) < 2:
        banner_row, starts = -1, [(0, "EMPLOYEE")]
    starts = [
        (_leading_product_start(header, col) if _banner_kind(text) == "product" else col, text)
        for col, text in starts
    ]

    tables = {"employee": _EMPLOYEE_ROLES, "dependant": _DEPENDANT_ROLES,
              "product": _PRODUCT_ROLES}
    blocks: list[ElBlock] = []
    for i, (col, text) in enumerate(starts):
        kind = _banner_kind(text)
        # Columns left of the first banner (movement code, staff ID) belong to it.
        start = 0 if i == 0 else col
        stop = starts[i + 1][0] if i + 1 < len(starts) else width
        columns = _columns(header, start, stop, tables[kind])
        hint = (
            _code_hint(text, [c.header for c in columns]) if kind == "product" else None
        )
        blocks.append(ElBlock(kind, text, tuple(columns), hint))

    trailing = tuple(
        ElColumn(c, get_column_letter(c + 1), norm_header(header[c]), role)
        for c in range(width)
        if norm_header(header[c])
        and (role := _role(norm_header(header[c]), _TRAILING_ROLES)) != "other"
    )
    layout = ElLayout(sheet_name, banner_row, header_row, tuple(blocks), trailing)
    employee = layout.employee
    dob = employee.column("date_of_birth") if employee else None
    reference = None
    if dob is not None:
        for r in range(min(banner_row if banner_row >= 0 else header_row, header_row)):
            value = rows[r][dob.index] if dob.index < len(rows[r]) else None
            if (found := as_date(value)) is not None:
                reference = found
    return ElLayout(sheet_name, banner_row, header_row, tuple(blocks), trailing, reference)
