"""A company-wide, source-aware employee listing, separate from submissions."""

from __future__ import annotations

import json
import re
from collections import defaultdict
from collections.abc import Iterator
from dataclasses import asdict
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.core.clock import today
from app.models import Employee, PolicyYear
from app.schemas.underwriting_reporting import REPORT_DETAIL_COLUMNS
from app.services.built_in_listings import _number
from app.services.full_el_data import FullELData, ProductProfile, load_full_el
from app.services.full_el_pricing import MemberLine, member_line
from app.services.insurer_listings import _ident, member_id_for_insurer
from app.services.insurer_reports import append_safe, as_date, last_day_of_service
from app.services.roster_attributes import (
    DEPENDANT_ID_KEYS,
    DOB_KEYS,
    EMPLOYEE_ID_KEYS,
    NAME_KEYS,
    REL_KEYS,
    first_value,
    mask_nric,
)

PALETTE = {"heading": "17365D", "header": "E7EEF6", "text": "FFFFFF", "review": "FFF2CC"}
SummaryGroups = dict[tuple[str, ...], dict[str, float]]
_BASE_FIELDS = [
    ("Entity", ("entity", "company", "subsidiary")),
    ("Gender", ("gender", "sex")),
    ("Marital Status", ("marital_status",)),
    ("Work Pass", ("pass", "pass_type", "work_pass")),
    ("Nationality", ("nationality",)),
    ("Country of Work", ("country_of_work", "country")),
    ("Employment Status", ("employment_status",)),
    ("Category", ("category",)),
    ("Designation", ("designation",)),
    ("Grade", ("job_grade", "grade")),
    ("Division", ("division",)),
    ("Department", ("department",)),
    ("Cost Centre", ("cost_centre", "cost_center")),
    ("Depot", ("depot",)),
    ("Location", ("location_description", "location")),
    ("Salary Currency", ("currency", "salary_currency")),
    ("Monthly Salary", ("salary", "monthly_salary", "basic_salary")),
    ("Hire Date", ("date_of_hire", "hire_date")),
    ("Confirmation Date", ("confirmation_date",)),
    ("Effective Date", ("effective_date",)),
    ("Email", ("email", "email_address")),
    ("Mobile", ("mobile", "mobile_phone")),
    ("Bank Code", ("bank_code",)),
    ("Branch Code", ("branch_code",)),
    ("Bank Account No.", ("bank_account_no",)),
    ("Retrenchment Start", ("retrenchment_start_date", "retrenchment_start")),
    ("Retrenchment End", ("retrenchment_end_date", "retrenchment_end")),
    ("Reemployment Start", ("reemployment_start_date", "reemployment_start")),
    ("Reemployment End", ("reemployment_end_date", "reemployment_end")),
    ("Remarks", ("remarks",)),
    ("Recorded Movement", ("movement", "movement_type", "action")),
    ("Change Notes", ("change_notes", "internal_remarks")),
]
_COVER_FIELDS = [
    ("Coverage Status", "status"),
    ("Category", "category"),
    ("Plan", "plan"),
    ("Family Group", "family"),
    ("Age", "age"),
    ("Cover Start", "start"),
    ("Cover End", "end"),
    ("Issued Policy No.", "policy_number"),
]
_SI_FIELDS = [
    ("Eligible SI", "eligible"),
    ("Guaranteed SI", "guaranteed"),
    ("Pending SI", "pending"),
    ("Accepted SI", "accepted"),
]
_UW_FIELDS = [
    ("UW Decision", "decision"),
    ("UW Decision Date", "decided_on"),
    ("UW Requirements", "requirements"),
    ("UW Remarks", "uw_remarks"),
]
_PRICE_FIELDS = [
    ("Rate Basis", "rate_basis"),
    ("Annual Rate", "rate"),
    ("Annual Net", "net"),
    ("Annual GST", "gst"),
    ("Annual Gross", "gross"),
    ("Pricing Status / Basis", "pricing"),
]


def _columns(profile: ProductProfile) -> list[tuple[str, str]]:
    return [
        *_COVER_FIELDS,
        *(_SI_FIELDS if profile.block.lump_sum else []),
        *_UW_FIELDS,
        *REPORT_DETAIL_COLUMNS,
        *_PRICE_FIELDS,
    ]


def _flat(value: Any, prefix: str = "") -> Iterator[tuple[str, Any]]:
    if isinstance(value, dict):
        for key, item in value.items():
            if key != "source_rows":
                yield from _flat(item, f"{prefix}.{key}" if prefix else str(key))
    elif isinstance(value, list):
        for index, item in enumerate(value, 1):
            yield from _flat(item, f"{prefix}[{index}]")
    elif value not in (None, ""):
        yield prefix, value


def _table(wb: Workbook, title: str, headers: list[str]) -> Worksheet:
    ws = wb.create_sheet(title)
    append_safe(ws, headers)
    return ws


def _format(ws: Worksheet, header_row: int = 1, freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.print_title_rows = f"1:{header_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    for cell in ws[header_row]:
        cell.font = Font(name="Arial", bold=True)
        cell.fill = PatternFill("solid", fgColor=PALETTE["header"])
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = 22
    ws.row_dimensions[header_row].height = 45
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "dd/mm/yyyy"
            elif isinstance(cell.value, float):
                cell.number_format = "#,##0.00####"
    if ws.title in ("Basis of Cover", "Source Setup", "Setup & Data Gaps", "Declaration"):
        for key in ("C", "D", "E"):
            ws.column_dimensions[key].width = 64
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def _member_sheet(
    wb: Workbook, title: str, headers: list[str], profiles: list[ProductProfile],
) -> Worksheet:
    ws = wb.create_sheet(title)
    groups = ["Member details"] * len(headers)
    columns = list(headers)
    spans = [(1, len(headers), "Member details")]
    for p in profiles:
        label = f"{p.block.report_code} | {p.block.insurer or 'Insurer not set'}"
        fields = _columns(p)
        start = len(columns) + 1
        columns.extend(f"{p.block.report_code} {name}" for name, _ in fields)
        columns.extend((f"{p.block.report_code} Administration", f"{p.block.report_code} Currency"))
        groups.extend([label] * (len(fields) + 2))
        spans.append((start, len(columns), label))
    append_safe(ws, groups)
    append_safe(ws, columns)
    for start, end, _ in spans:
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        cell = ws.cell(1, start)
        cell.fill = PatternFill("solid", fgColor=PALETTE["heading"])
        cell.font = Font(name="Arial", bold=True, color=PALETTE["text"])
    ws.row_dimensions[1].height = 24
    return ws


def _line_values(p: ProductProfile, line: MemberLine) -> list[Any]:
    values = []
    for _, key in _columns(p):
        value = getattr(line, key, line.report_details.get(key))
        values.append(as_date(value) if key.endswith("_date") else value)
    return [*values, p.rules.admin, p.rules.currency]


def _employee_values(emp: Employee, masked: bool) -> list[Any]:
    attrs = emp.attribute_values or {}
    values = [
        emp.staff_id,
        emp.employee_name or "",
        _ident(attrs, EMPLOYEE_ID_KEYS, masked),
        as_date(first_value(attrs, DOB_KEYS)),
        emp.status,
        last_day_of_service(emp),
    ]
    for name, keys in _BASE_FIELDS:
        value: Any = first_value(attrs, keys)
        if name == "Monthly Salary":
            value = _number(value)
        elif name.endswith(("Date", "Start", "End")):
            value = as_date(value)
        values.append(value)
    values.append(emp.updated_at.replace(tzinfo=None) if emp.updated_at else None)
    return values


def _basis(wb: Workbook, data: FullELData, employee_status: str, masked: bool) -> None:
    ws = _table(wb, "Basis of Cover", ["Product", "Section", "Field", "Value", "Source / Status"])
    meta = {
        "Company": data.company,
        "Benefit year": data.year.year,
        "Benefit period start": data.year.start_date,
        "Benefit period end": data.year.end_date,
        "Export date": today(),
        "Employee population": employee_status,
        "NRIC treatment": "Masked" if masked else "Unmasked",
        "Employees": len(data.employees),
        "Dependants on file": len(data.dependants),
        "Pricing": "Annual only. No proration. Unknown values are blank, never zero.",
        "Family premiums": "Flat / family-tier premiums appear once on the employee row.",
        "Summary scope": "Resolved cover overlapping this year, including in-period leavers.",
        "Movement": "Recorded movements only; no A/D/C comparison against an earlier file.",
        "Readiness": f"{len(data.gaps)} grouped setup / data gaps. Review before external use.",
    }
    for key, value in meta.items():
        append_safe(ws, ["", "Report", key, value, "Current company snapshot"])
    for p in data.profiles:
        code, term = p.block.product.code, p.term
        values = {
            "Insurer": p.block.insurer,
            "Setup status": p.setup.status if p.setup else "Missing",
            "Product cover start": term.coverage_start,
            "Product cover end": term.coverage_end,
            "Period inherited": term.is_default,
            **asdict(p.rules),
            "GST taxable": term.gst_included,
            "GST rate (%)": term.gst_rate,
            "Free cover limit": term.free_cover_limit,
            "NEL age (ANB)": term.nel_age_limit,
            "Underwriting required": term.underwriting_required,
            "Policy number assignments": json.dumps(term.policy_number_mappings or []),
            "Cover description": p.answers.get("cover_description", ""),
        }
        for key, value in values.items():
            append_safe(ws, [code, "Product", key, value, p.source])
        for cat in p.categories:
            facts = {
                "Description": cat.raw_description,
                "Review status": cat.status,
                "Mapping status": cat.rule_status,
                "Matching rule": cat.rule_human_readable,
                "Participation": cat.participation_detail,
                "Rates and cover": cat.plan_assignments,
            }
            for key, value in _flat(facts):
                append_safe(ws, [code, cat.display_name, key, value, cat.source_ref or ""])
    _format(ws)


def _source_setup(wb: Workbook, data: FullELData) -> None:
    ws = _table(
        wb, "Source Setup", ["Product", "Review Status", "Setup Field", "Stored Value", "Source"]
    )
    for p in data.profiles:
        for key, value in _flat(p.answers):
            append_safe(
                ws,
                [
                    p.block.product.code,
                    p.setup.status if p.setup else "Missing",
                    key,
                    value,
                    p.source,
                ],
            )
    _format(ws)


def _add_summary(
    groups: SummaryGroups, p: ProductProfile, emp: Employee, line: MemberLine, dependant: bool,
) -> None:
    if not line.included:
        return
    entity = first_value(emp.attribute_values or {}, ("entity", "company", "subsidiary")) or ""
    key = (
        p.block.product.code,
        p.block.insurer,
        entity,
        line.category,
        line.plan,
        line.family,
        p.rules.admin,
        p.rules.currency,
    )
    group = groups[key]
    group["dependants" if dependant else "employees"] += 1
    if line.premium_unit:
        group["units"] += 1
        for field in ("net", "gst", "gross"):
            value = getattr(line, field)
            if value is not None:
                group[field] += value
                group[f"{field}_known"] += 1


def _summary(wb: Workbook, groups: SummaryGroups) -> None:
    ws = _table(
        wb,
        "Headcount & Annual Premium",
        [
            "Product",
            "Insurer",
            "Entity",
            "Category",
            "Plan",
            "Family Group",
            "Administration",
            "Currency",
            "Employees",
            "Dependants",
            "Premium Units",
            "Unpriced Units",
            "Known Annual Net",
            "Known Annual GST",
            "Known Annual Gross",
            "Total Annual Gross",
            "Completeness",
        ],
    )
    for key, group in sorted(groups.items()):
        units = group["units"]
        complete = units > 0 and group["gross_known"] == units
        append_safe(
            ws,
            [
                *key,
                group["employees"],
                group["dependants"],
                units,
                units - group["net_known"],
                *[
                    round(group[field], 2) if group[f"{field}_known"] else None
                    for field in ("net", "gst", "gross")
                ],
                round(group["gross"], 2) if complete else None,
                "Complete" if complete else "Incomplete - known subtotals only",
            ],
        )
    _format(ws)


def _include_empty_products(groups: SummaryGroups, profiles: list[ProductProfile]) -> None:
    represented = {key[0] for key in groups}
    for p in profiles:
        if p.block.product.code not in represented:
            key = (
                p.block.product.code,
                p.block.insurer,
                "",
                "",
                "",
                "",
                p.rules.admin,
                p.rules.currency,
            )
            groups.setdefault(key, defaultdict(float))


def _gaps(wb: Workbook, data: FullELData) -> None:
    ws = _table(
        wb,
        "Setup & Data Gaps",
        [
            "Product",
            "Field",
            "Issue",
            "Action",
            "Source",
            "Occurrences",
            "Example Staff IDs (up to 10)",
        ],
    )
    for gap in data.gaps.values():
        append_safe(
            ws,
            [
                gap.product,
                gap.field,
                gap.issue,
                gap.action,
                gap.source,
                gap.count,
                ", ".join(gap.staff_ids),
            ],
        )
    if not data.gaps:
        append_safe(ws, ["", "", "No detected setup or data gaps", "", "", 0, ""])
    _format(ws)


def _declaration(wb: Workbook) -> None:
    ws = _table(wb, "Declaration", ["Item", "Completion", "Declaration / instruction"])
    for row in [
        [
            "Review",
            "Not signed",
            "This generated workbook does not record consent or a signed declaration.",
        ],
        [
            "Accuracy",
            "",
            "By submitting this file, the company declares the information is accurate and "
            "complete and will notify Inspro Insurance Brokers Pte. Ltd. of changes.",
        ],
        [
            "Consent",
            "",
            "The submitting company confirms it has obtained the necessary third-party "
            "consent to provide their data for the relevant purposes.",
        ],
        ["Authorised representative", "", "Complete before submission"],
        ["Designation", "", ""],
        ["Signature", "", ""],
        ["Date", "", ""],
    ]:
        append_safe(ws, row)
    _format(ws)


def build_full_el(
    db: Session, py: PolicyYear, *, masked: bool = True, employee_status: str = "all"
) -> Workbook:
    data = load_full_el(db, py, employee_status)
    wb = Workbook()
    wb.remove(wb.active)
    insurers = sorted({p.block.insurer for p in data.profiles if p.block.insurer})
    headers = [
        "Staff ID",
        "Employee Name",
        "Identification No.",
        "Date of Birth",
        "Employee Status",
        "Last Day of Service",
        *[name for name, _ in _BASE_FIELDS],
        "Last Updated",
        *[f"{insurer} Member ID" for insurer in insurers],
    ]
    employees = _member_sheet(wb, "Full EL", headers, data.profiles)
    groups: SummaryGroups = defaultdict(lambda: defaultdict(float))
    for emp in data.employees:
        row = _employee_values(emp, masked)
        row.extend(member_id_for_insurer(emp.attribute_values, insurer) for insurer in insurers)
        if not _ident(emp.attribute_values or {}, EMPLOYEE_ID_KEYS, False):
            data.gap(
                "",
                "Identification",
                "Employee national ID is missing",
                "Complete the roster ID.",
                staff_id=emp.staff_id,
            )
        for p in data.profiles:
            line = member_line(data, p, emp)
            row.extend(_line_values(p, line))
            _add_summary(groups, p, emp, line, False)
            if not line.included and line.status == "No resolved cover":
                data.gap(
                    p.block.product.code,
                    "Coverage",
                    line.status,
                    "Check eligibility, elections and dates; absence does not establish exclusion.",
                    p.source,
                    emp.staff_id,
                )
        append_safe(employees, row)
    _dependants(wb, data, groups, masked)
    _basis(wb, data, employee_status, masked)
    _source_setup(wb, data)
    _include_empty_products(groups, data.profiles)
    _summary(wb, groups)
    _gaps(wb, data)
    _declaration(wb)
    _format(employees, 2, "C3")
    wb.move_sheet("Basis of Cover", offset=-2)
    if masked:
        # Source wording and UW remarks can themselves quote a national ID.
        pattern = re.compile(r"(?<![A-Za-z0-9])[STFGM]\d{7}[A-Z](?![A-Za-z0-9])", re.I)
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str) and pattern.search(cell.value):
                        cell.value = pattern.sub(lambda match: mask_nric(match.group()), cell.value)
    return wb


def _dependants(wb: Workbook, data: FullELData, groups: SummaryGroups, masked: bool) -> None:
    ws = _member_sheet(
        wb,
        "Dependants",
        [
            "Staff ID",
            "Employee Name",
            "Employee Identification No.",
            "Dependant ID",
            "Dependant Name",
            "Dependant Identification No.",
            "Relationship",
            "Date of Birth",
            "Gender",
            "Status",
            "Effective Date",
            "Last Day of Cover",
            "Remarks",
        ],
        data.profiles,
    )
    employees = {e.id: e for e in data.employees}
    for dep in data.dependants:
        emp, attrs = employees.get(dep.employee_id or ""), dep.attribute_values or {}
        row = [
            emp.staff_id if emp else "",
            emp.employee_name if emp else "",
            _ident(emp.attribute_values or {}, EMPLOYEE_ID_KEYS, masked) if emp else "",
            dep.id,
            first_value(attrs, NAME_KEYS),
            _ident(attrs, DEPENDANT_ID_KEYS, masked),
            first_value(attrs, REL_KEYS),
            as_date(first_value(attrs, DOB_KEYS)),
            first_value(attrs, ("gender", "sex")),
            dep.status,
            as_date(attrs.get("effective_date")),
            last_day_of_service(dep),
            attrs.get("remarks"),
        ]
        if not emp:
            data.gap(
                "",
                "Dependant link",
                "Dependant has no linked employee",
                "Link the dependant before assigning coverage.",
            )
        for p in data.profiles:
            line = (
                member_line(data, p, emp, dep) if emp else MemberLine(status="Unlinked dependant")
            )
            row.extend(_line_values(p, line))
            if emp:
                _add_summary(groups, p, emp, line, True)
        append_safe(ws, row)
    _format(ws, 2, "E3")
