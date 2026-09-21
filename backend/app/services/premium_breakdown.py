"""Resolved member premiums and Flex funding, without inventing missing allocations."""

from __future__ import annotations

from collections import defaultdict
from decimal import Decimal

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Employee, PolicyYear, Product
from app.services.benefit_statement import build_benefit_statement
from app.services.insurer_reports import append_safe, autosize, bold_header
from app.services.product_insurer import insurer_map
from app.services.product_terms import gst_multiplier, resolve_terms


def build_premium_breakdown(db: Session, py: PolicyYear) -> Workbook:
    terms = {term.code: term for term in resolve_terms(db, py)}
    products = list(
        db.scalars(
            select(Product).where(
                Product.id.in_(
                    [term.product_id for term in terms.values()],
                )
            )
        )
    )
    insurers_by_id = insurer_map(db, py.id, products)
    insurers = {product.code: insurers_by_id.get(product.id, "") for product in products}
    workbook = Workbook()
    premiums = workbook.active
    premiums.title = "Member Premiums"
    append_safe(
        premiums,
        [
            "Staff ID",
            "Employee",
            "Entity",
            "Cost Centre",
            "Product",
            "Insurer",
            "Policy Number",
            "Plan",
            "Covered Dependants",
            "Coverage Start",
            "Coverage End",
            "Currency",
            "Annual Premium Excl GST",
            "GST",
            "Annual Premium Incl GST",
            "Rate Basis",
            "Data Status",
        ],
    )
    funding = workbook.create_sheet("Flex Funding")
    append_safe(
        funding,
        [
            "Staff ID",
            "Employee",
            "Cost Centre",
            "Currency",
            "Product",
            "Plan",
            "Total Flex Charge",
            "Dependant Portion",
            "Employee Coverage Portion",
            "Wallet Allowance",
            "Wallet Proration",
            "Leave Adjustment",
        ],
    )
    totals: dict[tuple[str, str, str], list[Decimal | int]] = defaultdict(
        lambda: [Decimal("0"), Decimal("0"), Decimal("0"), 0, 0],
    )
    employees = db.scalars(
        select(Employee)
        .where(
            Employee.client_id == py.client_id,
            Employee.policy_year_id == py.id,
        )
        .order_by(Employee.staff_id)
    ).all()
    for employee in employees:
        statement = build_benefit_statement(db, employee)
        attrs = employee.attribute_values or {}
        cost_centre = str(attrs.get("cost_centre") or attrs.get("cost_center") or "")
        for line in statement.coverage:
            term = terms.get(line.product_code)
            financials = line.financials
            gross = financials.annual_premium if financials else None
            multiplier = gst_multiplier(term.gst_included, term.gst_rate) if term else 1.0
            # Resolved financials already apply GST; never add it a second time.
            net = round(gross / multiplier, 2) if gross is not None else None
            tax = round(gross - net, 2) if gross is not None and net is not None else None
            insurer = insurers.get(line.product_code, "")
            append_safe(
                premiums,
                [
                    employee.staff_id,
                    employee.employee_name,
                    attrs.get("entity"),
                    cost_centre,
                    line.product_name or line.product_code,
                    insurer,
                    term.policy_number if term else None,
                    line.plan_code,
                    "; ".join(dep.name or "Unnamed dependant" for dep in line.covered_dependants),
                    term.coverage_start if term else py.start_date,
                    term.coverage_end if term else py.end_date,
                    "SGD",
                    net,
                    tax,
                    gross,
                    financials.rate_basis if financials else None,
                    "Resolved annual premium"
                    if gross is not None
                    else "Per-member premium unavailable",
                ],
            )
            bucket = totals[(cost_centre, insurer, line.product_name or line.product_code)]
            bucket[3] += 1
            if gross is None:
                bucket[4] += 1
            else:
                for i, amount in enumerate((net, tax, gross)):
                    bucket[i] += Decimal(str(amount or 0))
        flex = statement.flex
        if flex:
            for tag in flex.price_tag_lines:
                own = (
                    round(tag.price_tag - tag.dependant_tag, 2)
                    if tag.price_tag is not None and tag.dependant_tag is not None
                    else None
                )
                append_safe(
                    funding,
                    [
                        employee.staff_id,
                        employee.employee_name,
                        cost_centre,
                        flex.currency,
                        tag.product_code,
                        tag.plan_code,
                        tag.price_tag,
                        tag.dependant_tag,
                        own,
                        flex.wallet_amount,
                        flex.proration.note if flex.proration else "Full annual",
                        flex.leave_flex_amount,
                    ],
                )
    summary = workbook.create_sheet("Premium Summary")
    append_safe(
        summary,
        [
            "Cost Centre",
            "Insurer",
            "Product",
            "Currency",
            "Known Net Premium",
            "Known GST",
            "Known Gross Premium",
            "Coverage Lines",
            "Unpriced Lines",
        ],
    )
    for (cost_centre, insurer, product), amounts in sorted(totals.items()):
        append_safe(summary, [cost_centre, insurer, product, "SGD", *amounts])
    notes = workbook.create_sheet("Read Me")
    for row in [
        ["Scope", "Current resolved configuration for this benefit year; not an insurer invoice."],
        [
            "Currency",
            "Insurance premiums use policy currency SGD. Flex rows state their currency.",
        ],
        [
            "Unknown premiums",
            "Blank amounts are unknown, not free cover. Summary totals include known amounts only.",
        ],
        [
            "Dependants",
            "Names identify included lives. Family premiums are not divided between dependants.",
        ],
        [
            "Funding",
            "Flex charges are wallet debits, not premiums or personal payroll deductions.",
        ],
        [
            "Employer / employee",
            "Insurance funding shares are not recorded. Flex cost sharing does not imply them.",
        ],
        [
            "Adjustments",
            "Recorded Flex proration and leave adjustments are shown. "
            "Insurance endorsement/invoice reconciliation is not inferred.",
        ],
        [
            "Repeated wallet values",
            "Wallet and leave amounts repeat per product for context; do not sum these columns.",
        ],
    ]:
        append_safe(notes, row)
    for sheet in workbook:
        bold_header(sheet)
        autosize(sheet)
        sheet.freeze_panes = "A2"
        if sheet is not notes:
            sheet.auto_filter.ref = sheet.dimensions
    return workbook
