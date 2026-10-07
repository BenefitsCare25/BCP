"""Full Employee Listing: the broker-reviewed product setup drives the export.

GTL carries an applied reviewed cap on its product terms that disagrees with
the slip's stated maximum, and reviewed ALB ages to its own cover start. GPA
is an older confirmed setup without applied rules (its header cap is the
reviewed one) whose eligibility wording says ANB.
"""
from __future__ import annotations

import os
from pathlib import Path

import pytest

TEST_DB = Path(__file__).parent / "_test_full_el.db"
os.environ["INSPRO_DATABASE_URL"] = f"sqlite:///{TEST_DB}"

from datetime import date  # noqa: E402

from openpyxl.workbook import Workbook  # noqa: E402

from app.core.auth import DEMO_BROKER_FIRM_ID  # noqa: E402
from app.db.base import Base  # noqa: E402
from app.db.session import SessionLocal, engine  # noqa: E402
from app.models import (  # noqa: E402
    Category,
    Client,
    Employee,
    PolicyYear,
    Product,
    ProductSetup,
    ProductTerm,
)
from app.models.policy_year import PolicyYearStatus  # noqa: E402
from app.models.product_setup import ProductSetupStatus  # noqa: E402
from app.services.full_el import build_full_el  # noqa: E402
from app.services.full_el.context import load_context  # noqa: E402

CLIENT_ID = "00000000-0000-0000-0000-0000000f1000"
PY_ID = "00000000-0000-0000-0000-0000000f1001"
GTL_PROD = "00000000-0000-0000-0000-0000000f1010"
GPA_PROD = "00000000-0000-0000-0000-0000000f1011"
GTL_CAT = "00000000-0000-0000-0000-0000000f1020"
GPA_CAT = "00000000-0000-0000-0000-0000000f1021"
EMP_ID = "00000000-0000-0000-0000-0000000f1101"


def _slip_cap(amount: str) -> list[dict]:
    return [{"rows": [["Maximum sum insured", f"S${amount}"]]}]


@pytest.fixture(scope="module", autouse=True)
def _setup_db():
    if TEST_DB.exists():
        TEST_DB.unlink()
    Base.metadata.create_all(bind=engine)
    from scripts.seed_demo import seed
    seed()
    with SessionLocal() as s:
        s.add(Client(id=CLIENT_ID, name="Full EL Co", broker_firm_id=DEMO_BROKER_FIRM_ID))
        s.flush()
        s.add(PolicyYear(
            id=PY_ID, client_id=CLIENT_ID, year=2034,
            start_date=date(2034, 1, 1), end_date=date(2034, 12, 31),
            status=PolicyYearStatus.active,
        ))
        s.add_all([
            Product(id=GTL_PROD, client_id=CLIENT_ID, code="GTL",
                    display_name="Group Term Life", insurer="TestSure"),
            Product(id=GPA_PROD, client_id=CLIENT_ID, code="GPA",
                    display_name="Group Personal Accident", insurer="TestSure"),
        ])
        s.flush()
        s.add_all([
            Category(id=GTL_CAT, policy_year_id=PY_ID, product_id=GTL_PROD,
                     display_name="All staff", raw_description="All staff",
                     plan_assignments={"basis": "24x", "rate_basis": "per_1000_si",
                                       "premium_rate": 2.12},
                     source="manual", status="confirmed"),
            Category(id=GPA_CAT, policy_year_id=PY_ID, product_id=GPA_PROD,
                     display_name="All staff", raw_description="All staff",
                     plan_assignments={"basis": "400000", "rate_basis": "per_1000_si",
                                       "premium_rate": 0.5},
                     source="manual", status="confirmed"),
        ])
        s.add_all([
            ProductSetup(
                policy_year_id=PY_ID, product_code="GTL",
                status=ProductSetupStatus.confirmed,
                answers={
                    "header": {
                        "admin_basis": "Headcount basis",
                        "el_age_basis": "ALB",
                        "el_age_reference": "Product cover start",
                        "el_currency": "SGD",
                        "el_admin_resolution": "Named basis above the FCL",
                        # A later draft edit; the applied term cap still governs.
                        "el_max_sum_insured": "900000",
                    },
                    "sections": _slip_cap("1,600,000"),
                },
            ),
            ProductSetup(
                policy_year_id=PY_ID, product_code="GPA",
                status=ProductSetupStatus.confirmed,
                answers={
                    "header": {"el_max_sum_insured": "300000",
                               "el_premium_si_basis": "Accepted SI"},
                    "eligibility": {"eligibility": "Up to age 70 next birthday (ANB)"},
                    "sections": _slip_cap("1,000,000"),
                },
            ),
            ProductTerm(policy_year_id=PY_ID, product_id=GTL_PROD,
                        coverage_start=date(2034, 3, 1),
                        report_rules={"max_sum_insured": 500000}),
        ])
        s.add(Employee(
            id=EMP_ID, client_id=CLIENT_ID, policy_year_id=PY_ID,
            staff_id="FE-1", employee_name="Re Viewed",
            attribute_values={"date_of_birth": "1990-06-15", "salary": "30000"},
            derived_attribute_values={},
            matched_categories=[
                {"category_id": GTL_CAT, "product_code": "GTL", "method": "rule"},
                {"category_id": GPA_CAT, "product_code": "GPA", "method": "rule"},
            ],
            source="csv_import", status="active",
        ))
        s.commit()
    yield
    engine.dispose()
    if TEST_DB.exists():
        TEST_DB.unlink()


def _build() -> tuple[Workbook, dict[str, dict[str, str]], dict[str, str]]:
    """The workbook, each product block's role -> column letter, employee letters."""
    with SessionLocal() as s:
        year = s.get(PolicyYear, PY_ID)
        assert year is not None
        ctx = load_context(s, year)
        blocks = {
            ctx.block_products[i][0]: {c.role: c.letter for c in b.columns}
            for i, b in enumerate(ctx.layout.blocks) if b.kind == "product"
        }
        employee = ctx.layout.employee
        assert employee is not None
        return build_full_el(s, year), blocks, {c.role: c.letter for c in employee.columns}


def test_reviewed_caps_replace_the_slip_maximum() -> None:
    wb, blocks, emp = _build()
    ws = wb["Employee Listing"]
    # Applied term cap (500,000), not the slip's 1,600,000 or the draft header.
    assert ws[f"{blocks['GTL']['eligible_si']}4"].value == f"=MIN(24*{emp['salary']}4,500000)"
    # Confirmed setup without applied rules: the reviewed header cap.
    assert ws[f"{blocks['GPA']['eligible_si']}4"].value == 300000


def test_reviewed_age_basis_and_reference_are_per_product() -> None:
    wb, blocks, emp = _build()
    ws = wb["Employee Listing"]
    dob, age = emp["date_of_birth"], emp["age"]
    # Bases disagree across products, so the shared employee age stays ANB.
    assert ws[f"{age}3"].value == "Age (ANB)"
    assert ws[f"{age}4"].value == (
        f'=IF(OR({dob}4="",{dob}4>${dob}$1),"",DATEDIF({dob}4,${dob}$1,"y")+1)'
    )
    assert ws[f"{blocks['GPA']['age']}4"].value == f"={age}4"
    gtl_age = blocks["GTL"]["age"]
    assert ws[f"{gtl_age}3"].value == "Age (ALB)"
    ref = "DATE(2034,3,1)"
    assert ws[f"{gtl_age}4"].value == (
        f'=IF(OR({dob}4="",{dob}4>{ref}),"",DATEDIF({dob}4,{ref},"y"))'
    )


def test_summary_and_gaps_carry_the_reviewed_rules() -> None:
    wb, _blocks, _emp = _build()
    summary = wb["Summary - Basis of Cover"]
    rows = {
        r[1]: r for r in summary.iter_rows(min_row=5, values_only=True) if r[1]
    }
    gtl = next(r for banner, r in rows.items() if "(GTL)" in banner)
    assert gtl[5] == "Headcount basis\nNamed basis above the FCL"
    assert gtl[6] == "SGD"
    gaps = {(r[0], r[1]) for r in wb["Setup & Data Gaps"].iter_rows(min_row=2, values_only=True)}
    assert ("GPA", "Premium SI basis is Accepted SI") in gaps
    assert not any("not reviewed" in finding for _code, finding in gaps)
