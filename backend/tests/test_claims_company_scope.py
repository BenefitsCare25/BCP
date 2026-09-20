"""All-years browsing is tenant scoped; search and exports share predicates."""
from datetime import UTC, date, datetime
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.auth import (
    DEMO_BROKER_FIRM_ID,
    DEMO_CLIENT_ID,
    DEMO_USER_ID,
    CurrentUser,
    get_current_user,
)
from app.db.session import SessionLocal
from app.main import app
from app.models import Claim, Client, Dependant, Employee, Plan, PolicyYear, Product
from scripts.seed_demo import seed

CURRENT = "scope-current"
HISTORIC = "scope-historic"
FOREIGN = "scope-foreign"


@pytest.fixture(scope="module", autouse=True)
def setup_scope():
    seed()
    with SessionLocal() as db:
        db.add(Client(id=FOREIGN, name="Other company", broker_firm_id=DEMO_BROKER_FIRM_ID))
        db.flush()
        for year_id, year, client_id in (
            (CURRENT, 2030, DEMO_CLIENT_ID),
            (HISTORIC, 2029, DEMO_CLIENT_ID),
            (FOREIGN, 2030, FOREIGN),
        ):
            db.add(PolicyYear(id=year_id, client_id=client_id, year=year,
                              start_date=date(year, 1, 1), end_date=date(year, 12, 31)))
            db.flush()
            db.add(Employee(id=year_id, client_id=client_id, policy_year_id=year_id,
                            staff_id=year_id, employee_name="Parent Employee",
                            attribute_values={}, derived_attribute_values={}, source="csv_import"))
            db.flush()
            db.add(Dependant(id=year_id, client_id=client_id, policy_year_id=year_id,
                             employee_id=year_id, attribute_values={"full_name": "Jo_%Child"}))
            db.flush()
            db.add(Product(id=year_id, client_id=client_id, code=year_id,
                           display_name=year_id, insurer=f"Insurer {year_id}"))
            db.flush()
            db.add(Plan(product_id=year_id, policy_year_id=year_id, code="P1", display_name="P1"))
            for index in range(2):
                db.add(Claim(id=f"{year_id}-{index}", client_id=client_id,
                             policy_year_id=year_id, employee_id=year_id,
                             dependant_id=year_id if index == 0 else None,
                             claim_kind="insured", product_code="GHS", claim_type="Inpatient",
                             incurred_date=date(year, 6, 1), amount_claimed=100,
                             status="paid" if index == 0 else "submitted",
                             submitted_at=datetime(year, 6, 2, tzinfo=UTC)))
        db.commit()


@pytest.fixture
def broker():
    app.dependency_overrides[get_current_user] = lambda: CurrentUser(
        user_id=DEMO_USER_ID, broker_firm_id=DEMO_BROKER_FIRM_ID,
        client_id=DEMO_CLIENT_ID, role="broker_admin",
    )
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.pop(get_current_user, None)


def test_all_years_includes_processed_historic_but_never_other_company(broker):
    current = broker.get("/api/v1/claims", params={"policy_year_id": CURRENT}).json()
    assert current["total"] == 2
    all_years = broker.get("/api/v1/claims", params={"policy_year_id": CURRENT, "all_years": True})
    assert all_years.status_code == 200, all_years.text
    rows = all_years.json()["items"]
    assert {row["id"] for row in rows} == {
        f"{year}-{index}" for year in (CURRENT, HISTORIC) for index in range(2)
    }
    assert all(row["policy_year_label"] for row in rows)
    assert any(row["status"] == "paid" and row["policy_year_id"] == HISTORIC for row in rows)
    denied = broker.get("/api/v1/claims", params={"policy_year_id": FOREIGN, "all_years": True})
    assert denied.status_code == 404


@pytest.mark.parametrize("search", ["child", "o_%c", "%", "_"])
def test_dependant_search_literal_case_insensitive_and_paginates(broker, search):
    params = {"policy_year_id": CURRENT, "all_years": True, "search": search, "limit": 1}
    first = broker.get("/api/v1/claims", params=params)
    assert first.status_code == 200, first.text
    assert first.json()["total"] == 2
    second = broker.get("/api/v1/claims", params={**params, "offset": 1}).json()
    assert second["total"] == 2
    assert first.json()["items"][0]["id"] != second["items"][0]["id"]
    assert all(
        row["dependant_name"] == "Jo_%Child"
        for row in first.json()["items"] + second["items"]
    )


def test_filtered_export_matches_list_and_has_original_year(broker):
    params = {"policy_year_id": CURRENT, "all_years": True, "status": "paid",
              "incurred_from": "2029-01-01", "incurred_to": "2029-12-31", "search": "CHILD"}
    rows = broker.get("/api/v1/claims", params=params).json()["items"]
    assert [row["id"] for row in rows] == [f"{HISTORIC}-0"]
    exported = broker.get("/api/v1/claims/register", params=params)
    assert exported.status_code == 200, exported.text
    sheet = load_workbook(BytesIO(exported.content)).active
    values = list(sheet.values)
    assert len(values) == 2
    assert values[1][values[0].index("Claim ID")] == rows[0]["id"]
    assert "2029" in values[1][values[0].index("Benefit Year")]


def test_invalid_date_interval_is_rejected_on_list_and_export(broker):
    params = {"policy_year_id": CURRENT, "incurred_from": "2030-12-31", "incurred_to": "2030-01-01"}
    for path in ("/api/v1/claims", "/api/v1/claims/register"):
        assert broker.get(path, params=params).status_code == 422


def test_original_year_header_is_required_for_historic_detail(broker):
    claim_id = f"{HISTORIC}-0"
    wrong = broker.get(f"/api/v1/claims/{claim_id}", headers={"X-Inspro-Policy-Year-ID": CURRENT})
    assert wrong.status_code == 409
    assert "different benefit year" in wrong.json()["detail"]
    right = broker.get(f"/api/v1/claims/{claim_id}", headers={"X-Inspro-Policy-Year-ID": HISTORIC})
    assert right.status_code == 200, right.text
    assert right.json()["policy_year_id"] == HISTORIC


def test_explicit_draft_export_matches_filtered_list(broker):
    with SessionLocal() as db:
        db.add(Claim(id="scope-draft", client_id=DEMO_CLIENT_ID,
                     policy_year_id=CURRENT, employee_id=CURRENT, claim_kind="insured",
                     product_code="GHS", claim_type="Inpatient", incurred_date=date(2030, 5, 1),
                     amount_claimed=25, status="draft"))
        db.commit()
    try:
        params = {"policy_year_id": CURRENT, "status": "draft"}
        assert broker.get("/api/v1/claims", params=params).json()["total"] == 1
        response = broker.get("/api/v1/claims/register", params=params)
        assert response.status_code == 200, response.text
        assert list(load_workbook(BytesIO(response.content)).active.values)[1][0] == "scope-draft"
        legacy = broker.get("/api/v1/claims/register", params={"policy_year_id": CURRENT})
        assert len(list(load_workbook(BytesIO(legacy.content)).active.values)) == 3
    finally:
        with SessionLocal() as db:
            db.delete(db.get(Claim, "scope-draft"))
            db.commit()


def test_all_years_insurer_choices_include_history_and_exclude_other_company(broker):
    current = broker.get("/api/v1/claims/insurers", params={"policy_year_id": CURRENT})
    assert current.status_code == 200, current.text
    assert current.json() == [f"Insurer {CURRENT}"]
    historic = broker.get("/api/v1/claims/insurers",
                          params={"policy_year_id": CURRENT, "all_years": True})
    assert historic.status_code == 200, historic.text
    assert f"Insurer {HISTORIC}" in historic.json()
    assert f"Insurer {FOREIGN}" not in historic.json()
    denied = broker.get("/api/v1/claims/insurers", params={"policy_year_id": FOREIGN})
    assert denied.status_code == 404
