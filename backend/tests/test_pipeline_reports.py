from datetime import date
from io import BytesIO
from time import perf_counter
from types import SimpleNamespace as NS
from unittest.mock import MagicMock

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.core.auth import (
    DEMO_BROKER_FIRM_ID,
    DEMO_CLIENT_ID,
    DEMO_USER_ID,
    CurrentUser,
    Role,
    get_current_user,
)
from app.db.session import SessionLocal
from app.main import app
from app.models import AuditLog, BrokerFirm, Claim, Client, PolicyYear
from app.services import claim_placement, premium_breakdown

FOREIGN_FIRM_ID = "00000000-0000-0000-0000-0000000f1600"
FOREIGN_CLIENT_ID = "00000000-0000-0000-0000-0000000f1601"
FOREIGN_YEAR_ID = "00000000-0000-0000-0000-0000000f1602"


def _user(role: Role = "broker_admin") -> CurrentUser:
    return CurrentUser(
        user_id=DEMO_USER_ID,
        broker_firm_id=DEMO_BROKER_FIRM_ID,
        client_id=DEMO_CLIENT_ID,
        role=role,
    )


@pytest.fixture(scope="module", autouse=True)
def _report_scope() -> None:
    from scripts.seed_demo import seed

    seed()
    with SessionLocal() as db:
        db.add(BrokerFirm(id=FOREIGN_FIRM_ID, name="Foreign report firm"))
        db.flush()
        db.add(
            Client(
                id=FOREIGN_CLIENT_ID,
                name="Foreign report company",
                broker_firm_id=FOREIGN_FIRM_ID,
            )
        )
        db.flush()
        db.add(
            PolicyYear(
                id=FOREIGN_YEAR_ID,
                client_id=FOREIGN_CLIENT_ID,
                year=2035,
                start_date=date(2035, 1, 1),
                end_date=date(2035, 12, 31),
            )
        )
        db.commit()


@pytest.fixture
def demo_year_id() -> str:
    with SessionLocal() as db:
        year = db.query(PolicyYear).filter(PolicyYear.client_id == DEMO_CLIENT_ID).first()
        assert year is not None
        return year.id


def test_filing_snapshot_survives_policy_number_edit_and_preserves_other_metadata(monkeypatch):
    claim = Claim(claim_kind="insured", product_code="GHS", intake_meta={"received_via": "email"})
    current = {
        "product_code": "GHS",
        "policy_number": "OLD",
        "insurer": "A",
        "product_name": "Hospital",
    }
    monkeypatch.setattr(claim_placement, "current_placement", lambda *_: dict(current))
    claim_placement.capture_claim_placement(None, claim)
    current["policy_number"] = "NEW"
    claim_placement.capture_claim_placement(None, claim)
    assert claim_placement.placement_cells(None, claim) == ["OLD", "A", "Hospital", "At filing"]
    assert claim.intake_meta["received_via"] == "email"
    legacy = Claim(claim_kind="insured", product_code="GHS")
    assert claim_placement.placement_cells(None, legacy)[0] == "NEW"
    assert "no filing snapshot" in claim_placement.placement_cells(None, legacy)[3]


def test_changed_product_does_not_use_previous_product_snapshot(monkeypatch):
    claim = Claim(
        claim_kind="insured",
        product_code="GP",
        intake_meta={
            "placement_snapshot": {"product_code": "GHS", "policy_number": "HOSPITAL"},
        },
    )
    monkeypatch.setattr(
        claim_placement,
        "current_placement",
        lambda *_: {"product_code": "GP", "policy_number": "CLINIC"},
    )
    assert claim_placement.placement_cells(None, claim)[0] == "CLINIC"


def test_legacy_resubmission_does_not_invent_filing_snapshot(monkeypatch):
    claim = Claim(claim_kind="insured", product_code="GHS", status="needs_info")
    monkeypatch.setattr(claim_placement, "current_placement", lambda *_: {"policy_number": "NEW"})
    claim_placement.capture_claim_placement(None, claim)
    assert claim.intake_meta is None
    assert (
        claim_placement.placement_cells(None, claim)[-1]
        == "Current configuration; no filing snapshot"
    )


def test_premium_breakdown_gst_unknown_prices_and_formula_safety(monkeypatch):
    year = PolicyYear(
        id="year",
        client_id="client",
        year=2026,
        start_date=date(2026, 1, 1),
        end_date=date(2026, 12, 31),
    )
    term = NS(
        code="GHS",
        product_id="product",
        policy_number="P123",
        gst_included=True,
        gst_rate=9,
        coverage_start=year.start_date,
        coverage_end=year.end_date,
    )
    employee = NS(
        id="ee",
        staff_id="=evil()",
        employee_name="+evil()",
        attribute_values={"cost_centre": "Operations"},
    )
    db = MagicMock()
    member_result = MagicMock()
    member_result.all.return_value = [employee]
    db.scalars.side_effect = [[NS(id="product", code="GHS")], member_result]
    monkeypatch.setattr(premium_breakdown, "resolve_terms", lambda *_: [term])
    monkeypatch.setattr(premium_breakdown, "insurer_map", lambda *_: {"product": "Insurer"})
    coverage = [
        NS(
            product_code="GHS",
            product_name="Hospital",
            plan_code="A",
            covered_dependants=[],
            financials=NS(annual_premium=109, rate_basis="fixed"),
        ),
        NS(
            product_code="GP",
            product_name="Clinic",
            plan_code="B",
            covered_dependants=[],
            financials=None,
        ),
    ]
    flex = NS(
        currency="SGD",
        price_tag_lines=[
            NS(
                product_code="GHS",
                plan_code="A",
                price_tag=200,
                dependant_tag=50,
            )
        ],
        wallet_amount=500,
        proration=NS(note="6/12 months"),
        leave_flex_amount=-25,
    )
    monkeypatch.setattr(
        premium_breakdown,
        "build_benefit_statement",
        lambda *_: NS(coverage=coverage, flex=flex),
    )
    workbook = premium_breakdown.build_premium_breakdown(db, year)
    rows = list(workbook["Member Premiums"].values)
    assert rows[1][0] == "'=evil()"
    assert rows[1][1] == "'+evil()"
    assert rows[1][12:15] == (100, 9, 109)
    assert rows[2][12:15] == (None, None, None)
    assert rows[2][16] == "Per-member premium unavailable"
    funding = list(workbook["Flex Funding"].values)
    assert funding[1][6:12] == (200, 50, 150, 500, "6/12 months", -25)
    assert workbook["Premium Summary"].max_row == 3
    assert "Read Me" in workbook.sheetnames


def test_premium_breakdown_endpoint_authorization_and_audit(
    demo_year_id: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = Workbook()
    workbook.active.append(["Staff ID", "Annual Premium Incl GST"])
    workbook.active.append(["SAFE-1", 109])
    monkeypatch.setattr(
        premium_breakdown,
        "build_premium_breakdown",
        lambda *_: workbook,
    )
    path = f"/api/v1/policy-years/{demo_year_id}/reports/premium-breakdown"
    try:
        app.dependency_overrides[get_current_user] = lambda: _user("client_hr")
        with TestClient(app) as client:
            denied = client.get(path)
        assert denied.status_code == 403

        app.dependency_overrides[get_current_user] = lambda: _user()
        with TestClient(app) as client:
            concealed = client.get(
                f"/api/v1/policy-years/{FOREIGN_YEAR_ID}/reports/premium-breakdown"
            )
        assert concealed.status_code == 404

        app.dependency_overrides[get_current_user] = lambda: _user("broker_viewer")
        with TestClient(app) as client:
            viewer_download = client.get(path)
        assert viewer_download.status_code == 200

        app.dependency_overrides[get_current_user] = lambda: _user()
        with TestClient(app) as client:
            response = client.get(path)
        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert response.headers["content-disposition"].endswith('.xlsx"')
        downloaded = load_workbook(BytesIO(response.content), read_only=True)
        assert downloaded.active["A2"].value == "SAFE-1"

        with SessionLocal() as db:
            audit = (
                db.query(AuditLog)
                .filter(
                    AuditLog.action == "report.premium_breakdown",
                    AuditLog.entity_id == demo_year_id,
                    AuditLog.user_id == DEMO_USER_ID,
                )
                .order_by(AuditLog.created_at.desc())
                .first()
            )
            assert audit is not None
            assert audit.client_id == DEMO_CLIENT_ID
            assert audit.after == {"report": "premium-breakdown", "format": "xlsx"}
            assert audit.request_id
    finally:
        app.dependency_overrides.pop(get_current_user, None)


def test_premium_breakdown_1000_member_workbook_baseline(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    year = PolicyYear(
        id="year",
        client_id="client",
        year=2026,
        start_date=date(2026, 1, 1),
        end_date=date(2026, 12, 31),
    )
    employees = [
        NS(
            id=f"employee-{index}",
            staff_id=f"S{index:04d}",
            employee_name=f"Synthetic Employee {index}",
            attribute_values={"cost_centre": f"CC-{index % 10}"},
        )
        for index in range(1_000)
    ]
    employee_result = MagicMock()
    employee_result.all.return_value = employees
    db = MagicMock()
    db.scalars.side_effect = [[], employee_result]
    coverage = [
        NS(
            product_code="GHS",
            product_name="Hospital",
            plan_code="CORE",
            covered_dependants=[],
            financials=NS(annual_premium=109, rate_basis="flat"),
        )
    ]
    monkeypatch.setattr(premium_breakdown, "resolve_terms", lambda *_: [])
    monkeypatch.setattr(premium_breakdown, "insurer_map", lambda *_: {})
    monkeypatch.setattr(
        premium_breakdown,
        "build_benefit_statement",
        lambda *_: NS(coverage=coverage, flex=None),
    )

    started = perf_counter()
    workbook = premium_breakdown.build_premium_breakdown(db, year)
    elapsed = perf_counter() - started

    assert workbook["Member Premiums"].max_row == 1_001
    print(f"premium workbook 1,000-member baseline: {elapsed:.3f}s")
    assert elapsed < 5.0, f"1,000-member workbook took {elapsed:.3f}s"


def test_premium_breakdown_endpoint_100_request_latency_baseline(
    demo_year_id: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = Workbook()
    workbook.active.append(["Staff ID", "Annual Premium Incl GST"])
    workbook.active.append(["SAFE-1", 109])
    monkeypatch.setattr(
        premium_breakdown,
        "build_premium_breakdown",
        lambda *_: workbook,
    )
    path = f"/api/v1/policy-years/{demo_year_id}/reports/premium-breakdown"
    latencies: list[float] = []
    try:
        app.dependency_overrides[get_current_user] = lambda: _user()
        with TestClient(app) as client:
            for _ in range(100):
                started = perf_counter()
                response = client.get(path)
                latencies.append(perf_counter() - started)
                assert response.status_code == 200
    finally:
        app.dependency_overrides.pop(get_current_user, None)

    ordered = sorted(latencies)
    p50 = ordered[49]
    p95 = ordered[94]
    p99 = ordered[98]
    print(
        "premium endpoint 100-request baseline: "
        f"p50={p50:.3f}s p95={p95:.3f}s p99={p99:.3f}s"
    )
    assert p95 < 0.25, f"Premium download endpoint p95 was {p95:.3f}s"
