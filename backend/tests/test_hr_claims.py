"""Registration and response-shape regressions for delegated HR claims, and the
HR enrolment-form exports (masking, the admin-only bulk ZIP, rate limits)."""

from __future__ import annotations

from collections.abc import Iterator
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from typing import Any, cast

import pytest
from fastapi import HTTPException
from fastapi.routing import APIRoute
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from starlette.requests import Request

from app.api.v1 import hr_claims
from app.core.auth import CurrentUser, Role
from app.core.deps import require_write_access
from app.core.hr_auth import get_current_hr_user
from app.core.rate_limit import limiter
from app.db.session import SessionLocal
from app.main import app
from app.models import (
    AuditLog,
    BrokerFirm,
    Client,
    Employee,
    PolicyYear,
    User,
    UserClientAccess,
)
from app.models.claim import ORIGIN_HR
from app.models.enrollment_form import EnrollmentFormSubmission
from app.models.policy_year import PolicyYearStatus

FORMS_FIRM = "hr-forms-firm"
FORMS_CLIENT = "hr-forms-client"
FORMS_YEAR = "hr-forms-year"
FORMS_EMPLOYEE = "hr-forms-employee"
FORMS_SUBMISSION = "hr-forms-submission"
FORMS_HR = "hr-forms-user-hr"
FORMS_ADMIN = "hr-forms-user-admin"
EXPORTS = "/api/v1/hr/enrollment-forms"


def test_hr_claims_router_is_registered_outside_broker_write_gate() -> None:
    route = next(
        route
        for route in app.routes
        if isinstance(route, APIRoute)
        and route.path == "/api/v1/hr/claims"
        and "GET" in route.methods
    )

    dependency_calls = {dependency.call for dependency in route.dependant.dependencies}
    assert require_write_access not in dependency_calls


def test_claim_output_uses_document_display_name(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    claim = SimpleNamespace(
        id="claim-1",
        employee_id="employee-1",
        policy_year_id="year-1",
        reference_no="CLM-1",
        claim_kind="insured",
        claim_type="Outpatient",
        provider_name="Demo Clinic",
        invoice_number="INV-1",
        status="draft",
        incurred_date=None,
        amount_claimed=None,
        currency="SGD",
        created_by_user_id="user-1",
        intake_meta={
            "delegated_by_name": "Demo HR",
            "delegated_by_email": "hr@demo.test",
        },
        created_at=None,
        submitted_at=None,
    )
    document_slot = SimpleNamespace(
        key="invoice",
        display="Medical invoice",
        instructions="Upload the itemised invoice.",
    )
    monkeypatch.setattr(
        hr_claims,
        "setup_for_claim",
        lambda _db, _claim: SimpleNamespace(documents=[document_slot]),
    )
    monkeypatch.setattr(hr_claims, "claim_documents", lambda _db, _claim: [])

    class FakeSession:
        def get(self, _model: type[Any], _entity_id: str) -> SimpleNamespace:
            return SimpleNamespace(employee_name="Demo Employee")

    output = hr_claims._out(FakeSession(), claim)  # type: ignore[arg-type]

    assert output["doc_slots"] == [
        {
            "key": "invoice",
            "label": "Medical invoice",
            "instructions": "Upload the itemised invoice.",
        }
    ]
    assert output["submission_channel"] == "hr"
    assert output["submitted_by_name"] == "Demo HR"
    assert output["submitted_by_email"] == "hr@demo.test"
    assert output["can_add_evidence"] is True
    assert output["can_submit"] is True


def test_draft_records_hr_origin_and_submitter_snapshot(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    body = hr_claims.DelegatedClaimIn(
        employee_id="employee-1",
        claim_kind="flex",
        flex_category_name="Wellness",
        claim_type="Wellness",
        incurred_date="2026-09-20",
        provider_name="Demo Provider",
        invoice_number="HR-DEMO-1",
        amount_claimed=25,
        currency="SGD",
    )
    employee = SimpleNamespace(id="employee-1", client_id="client-1")
    claim = SimpleNamespace(id="claim-1", intake_meta=None)

    monkeypatch.setattr(hr_claims, "_employee", lambda *_args: employee)
    monkeypatch.setattr(hr_claims, "lock_claim_command_key", lambda *_args: None)
    monkeypatch.setattr(hr_claims, "replayed_claim_for_command", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(hr_claims, "create_claim", lambda *_args, **_kwargs: claim)
    monkeypatch.setattr(hr_claims, "write_audit", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(hr_claims, "record_claim_command", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(hr_claims, "_out", lambda _db, value: {"id": value.id})

    class FakeSession:
        def get(self, _model: type[Any], _entity_id: str) -> SimpleNamespace:
            return SimpleNamespace(display_name="Demo HR")

        def commit(self) -> None:
            pass

    user = CurrentUser(
        user_id="user-1",
        broker_firm_id="firm-1",
        client_id="client-1",
        role="client_hr",
        email="hr@demo.test",
    )
    request = Request({"type": "http", "method": "POST", "path": "/"})

    result = hr_claims.draft(request, body, "command-1", user, FakeSession())  # type: ignore[arg-type]

    assert result == {"id": "claim-1"}
    assert claim.origin == ORIGIN_HR
    assert claim.created_by_user_id == "user-1"
    assert claim.intake_meta == {
        "submission_channel": "hr",
        "delegated_by_user_id": "user-1",
        "delegated_by_name": "Demo HR",
        "delegated_by_email": "hr@demo.test",
    }


def test_hr_evidence_window_is_status_gated_not_portal_origin_gated() -> None:
    hr_claims._assert_evidence_mutable(SimpleNamespace(status="draft"))  # type: ignore[arg-type]
    hr_claims._assert_evidence_mutable(SimpleNamespace(status="needs_info"))  # type: ignore[arg-type]

    with pytest.raises(HTTPException) as decided:
        hr_claims._assert_evidence_mutable(SimpleNamespace(status="approved"))  # type: ignore[arg-type]

    assert decided.value.status_code == 403
    assert decided.value.detail == "Evidence is retained after a decision."


# ── HR enrolment-form exports ────────────────────────────────────────────────


def _hr_principal(user_id: str, role: str) -> CurrentUser:
    return CurrentUser(
        user_id=user_id,
        broker_firm_id=FORMS_FIRM,
        client_id=FORMS_CLIENT,
        role=cast(Role, role),
        email=f"{user_id}@example.test",
    )


def _cleanup_forms() -> None:
    with SessionLocal() as db:
        db.query(AuditLog).filter(AuditLog.client_id == FORMS_CLIENT).delete(
            synchronize_session=False
        )
        db.query(EnrollmentFormSubmission).filter(
            EnrollmentFormSubmission.id == FORMS_SUBMISSION
        ).delete(synchronize_session=False)
        db.query(UserClientAccess).filter(
            UserClientAccess.user_id.in_((FORMS_HR, FORMS_ADMIN))
        ).delete(synchronize_session=False)
        db.query(Employee).filter(Employee.id == FORMS_EMPLOYEE).delete(
            synchronize_session=False
        )
        db.query(PolicyYear).filter(PolicyYear.id == FORMS_YEAR).delete(
            synchronize_session=False
        )
        db.query(User).filter(User.id.in_((FORMS_HR, FORMS_ADMIN))).delete(
            synchronize_session=False
        )
        db.query(Client).filter(Client.id == FORMS_CLIENT).delete(synchronize_session=False)
        db.query(BrokerFirm).filter(BrokerFirm.id == FORMS_FIRM).delete(
            synchronize_session=False
        )
        db.commit()


@pytest.fixture
def hr_forms() -> Iterator[tuple[TestClient, dict[str, CurrentUser]]]:
    _cleanup_forms()
    with SessionLocal() as db:
        db.add(BrokerFirm(id=FORMS_FIRM, name="HR forms firm"))
        db.add(Client(id=FORMS_CLIENT, name="HR forms client", broker_firm_id=FORMS_FIRM))
        db.flush()
        db.add(PolicyYear(
            id=FORMS_YEAR, client_id=FORMS_CLIENT, year=2026,
            start_date=date(2026, 1, 1), end_date=date(2026, 12, 31),
            status=PolicyYearStatus.active,
        ))
        db.flush()
        db.add(Employee(
            id=FORMS_EMPLOYEE, client_id=FORMS_CLIENT, policy_year_id=FORMS_YEAR,
            staff_id="F-001", employee_name="Form Filer",
            attribute_values={"id_no": "S1234567D"},
        ))
        for user_id, role in ((FORMS_HR, "client_hr"), (FORMS_ADMIN, "client_admin")):
            db.add(User(
                id=user_id, email=f"{user_id}@example.test", display_name=user_id,
                broker_firm_id=FORMS_FIRM, role=role,
            ))
            db.add(UserClientAccess(user_id=user_id, client_id=FORMS_CLIENT))
        db.flush()
        db.add(EnrollmentFormSubmission(
            id=FORMS_SUBMISSION, client_id=FORMS_CLIENT, policy_year_id=FORMS_YEAR,
            employee_id=FORMS_EMPLOYEE, reference_no="EF-2026-90001",
            snapshot={
                "particulars": {"id_no": "S1234567D", "contact_no": "+6591234567"},
                "dependants": [
                    {"relationship": "Spouse", "name": "Kim Filer", "id_no": "T7654321Z"},
                ],
            },
        ))
        db.commit()

    active = {"user": _hr_principal(FORMS_HR, "client_hr")}
    app.dependency_overrides[get_current_hr_user] = lambda: active["user"]
    try:
        with TestClient(app) as client:
            yield client, active
    finally:
        app.dependency_overrides.pop(get_current_hr_user, None)
        _cleanup_forms()


def test_hr_summary_export_is_masked_and_audited_with_counts(
    hr_forms: tuple[TestClient, dict[str, CurrentUser]],
) -> None:
    client, _ = hr_forms

    res = client.get(f"{EXPORTS}/export.xlsx")

    assert res.status_code == 200, res.text
    workbook = load_workbook(BytesIO(res.content))
    summary = list(workbook["Enrolment forms"].iter_rows(values_only=True))
    family = list(workbook["Family members"].iter_rows(values_only=True))
    assert summary[1][summary[0].index("NRIC / FIN")] == "S******7D"
    assert family[1][family[0].index("NRIC / BC / FIN")] == "T******1Z"
    with SessionLocal() as db:
        row = db.query(AuditLog).filter(
            AuditLog.client_id == FORMS_CLIENT,
            AuditLog.action == "enrollment_form.export_xlsx",
        ).one()
    assert row.after == {"masked": True, "forms": 1, "family_members": 1}
    assert row.ip_address == "testclient"


def test_bulk_signed_form_zip_is_for_hr_administrators_only(
    hr_forms: tuple[TestClient, dict[str, CurrentUser]],
) -> None:
    client, active = hr_forms

    refused = client.get(f"{EXPORTS}/export.zip")
    active["user"] = _hr_principal(FORMS_ADMIN, "client_admin")
    allowed = client.get(f"{EXPORTS}/export.zip")

    assert refused.status_code == 403
    # Past the role gate: nothing to zip because no PDF was ever stored.
    assert allowed.status_code == 404
    assert allowed.json()["detail"] == "No forms match these filters."


def test_hr_bulk_exports_are_rate_limited(
    hr_forms: tuple[TestClient, dict[str, CurrentUser]],
) -> None:
    client, _ = hr_forms
    limiter.reset()
    limiter.enabled = True
    try:
        statuses = [client.get(f"{EXPORTS}/export.xlsx").status_code for _ in range(6)]
    finally:
        limiter.enabled = False
        limiter.reset()

    assert statuses == [200] * 5 + [429]
