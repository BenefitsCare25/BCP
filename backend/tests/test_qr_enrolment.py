"""No-login QR enrolment: shared employee form, broker review, HR link + PIN.

Synthetic company and people only. Storage and the database are the suite's
disposable temp stores (conftest); nothing touches backend/inspro.db.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import replace
from datetime import UTC, date, datetime, timedelta
from typing import Any
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.auth import DEMO_BROKER_FIRM_ID, CurrentUser, get_current_user
from app.db.session import SessionLocal
from app.main import app
from app.models import (
    AuditLog,
    Category,
    Client,
    Dependant,
    Employee,
    EnrollmentFormSubmission,
    EnrollmentWindow,
    HrDownloadBatch,
    HrDownloadGrant,
    Plan,
    PolicyYear,
    Product,
    PublicIntake,
    StoredDocument,
)
from app.models.category import CategoryStatus, SourceKind
from app.models.policy_year import PolicyYearStatus
from app.services.qr_enrolment import batches as batch_svc
from app.services.roster_attributes import sg_nric_check_letter
from scripts.seed_demo import seed

CLIENT_ID = "00000000-0000-0000-0000-0000000qr000"
OTHER_CLIENT = "00000000-0000-0000-0000-0000000qr0ff"
PY_ID = "00000000-0000-0000-0000-0000000qr001"
GHS_ID = "00000000-0000-0000-0000-0000000qr002"
GMM_ID = "00000000-0000-0000-0000-0000000qr003"
EMP1 = "00000000-0000-0000-0000-0000000qr010"


def _nric(prefix: str, digits: str) -> str:
    letter = sg_nric_check_letter(prefix, digits)
    assert letter is not None
    return f"{prefix}{digits}{letter}"


EMP_NRIC = _nric("S", "8012345")
SPOUSE_NRIC = _nric("S", "8212345")
CHILD_NRIC = _nric("T", "1512345")


def _broker(role: str = "broker_admin") -> CurrentUser:
    return CurrentUser(
        user_id="00000000-0000-0000-0000-0000000qrfff",
        broker_firm_id=DEMO_BROKER_FIRM_ID,
        client_id=CLIENT_ID,
        role=role,
    )


def _tiers(eo: float, es: float, ec: float, ef: float) -> dict[str, dict[str, float]]:
    return {k: {"rate": v} for k, v in (("EO", eo), ("ES", es), ("EC", ec), ("EF", ef))}


def _category(
    cid: str,
    product: str,
    name: str,
    plan: str,
    mode: dict[str, str],
    rates: dict[str, dict[str, float]],
) -> Category:
    return Category(
        id=cid,
        policy_year_id=PY_ID,
        product_id=product,
        priority=1,
        display_name=name,
        raw_description=name,
        plan_assignments={"plan_code": plan, "rate_basis": "tiered", "rate_tiers": rates},
        participation_detail=mode,
        source=SourceKind.system_generated.value,
        status=CategoryStatus.confirmed.value,
        human_modified=False,
    )


@pytest.fixture(scope="module", autouse=True)
def _setup_db():
    seed()
    with SessionLocal() as s:
        s.add(
            Client(
                id=CLIENT_ID,
                name="QR Test Co",
                legal_name="QR Test Company Pte Ltd",
                broker_firm_id=DEMO_BROKER_FIRM_ID,
                slug="qrco",
            )
        )
        s.add(
            Client(
                id=OTHER_CLIENT, name="Other Co", broker_firm_id=DEMO_BROKER_FIRM_ID, slug="otherco"
            )
        )
        s.flush()
        s.add(
            PolicyYear(
                id=PY_ID,
                client_id=CLIENT_ID,
                year=2029,
                start_date=date(2029, 4, 1),
                end_date=date(2030, 3, 31),
                status=PolicyYearStatus.active,
            )
        )
        s.flush()
        s.add(
            Product(
                id=GHS_ID,
                client_id=CLIENT_ID,
                code="GHS",
                display_name="Hospital & Surgical",
                insurer="Synthetic Insurer",
                has_dependants=True,
            )
        )
        s.add(
            Product(
                id=GMM_ID,
                client_id=CLIENT_ID,
                code="GMM",
                display_name="Major Medical",
                insurer="Synthetic Insurer",
                has_dependants=True,
            )
        )
        s.flush()
        for pid, code in ((GHS_ID, "P1"), (GHS_ID, "P2"), (GMM_ID, "M1")):
            s.add(
                Plan(
                    id=f"qr-plan-{code}",
                    product_id=pid,
                    policy_year_id=PY_ID,
                    code=code,
                    display_name=f"Plan {code}",
                    status="confirmed",
                )
            )
        compulsory = {"employee": "compulsory", "dependant": "voluntary"}
        voluntary = {"employee": "voluntary", "dependant": "voluntary"}
        s.add(
            _category(
                "qr-cat-senior",
                GHS_ID,
                "Senior management",
                "P1",
                compulsory,
                _tiers(400, 1000, 900, 1500),
            )
        )
        s.add(
            _category(
                "qr-cat-staff",
                GHS_ID,
                "All other staff",
                "P2",
                compulsory,
                _tiers(300, 800, 700, 1200),
            )
        )
        s.add(
            _category(
                "qr-cat-gmm", GMM_ID, "All employees", "M1", voluntary, _tiers(120, 270, 260, 420)
            )
        )
        s.add(
            Employee(
                id=EMP1,
                client_id=CLIENT_ID,
                policy_year_id=PY_ID,
                staff_id="Q-1",
                employee_name="Tan Ah Kow",
                attribute_values={},
                derived_attribute_values={},
                national_id_normalized=EMP_NRIC,
                source="csv_import",
                status="active",
            )
        )
        s.commit()
    yield


@pytest.fixture
def broker() -> TestClient:
    app.dependency_overrides[get_current_user] = _broker
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.pop(get_current_user, None)


@pytest.fixture(autouse=True)
def _clean():
    yield
    with SessionLocal() as s:
        for model in (HrDownloadGrant, HrDownloadBatch, PublicIntake, EnrollmentFormSubmission):
            s.query(model).delete()
        s.query(StoredDocument).delete()
        s.query(AuditLog).delete()
        s.execute(EnrollmentWindow.__table__.delete())
        s.commit()


def _window(broker: TestClient) -> str:
    wid = broker.post(
        f"/api/v1/policy-years/{PY_ID}/enrollment-windows",
        json={
            "name": "Annual enrolment",
            "window_type": "open",
            "opens_at": "2020-01-01T00:00:00Z",
            "closes_at": "2035-01-01T00:00:00Z",
        },
    ).json()["id"]
    assert broker.post(f"/api/v1/enrollment-windows/{wid}/open").status_code == 200
    config = broker.get(f"/api/v1/enrollment-windows/{wid}/form-config").json()["settings"]
    config["contributions"] = {
        "GHS": {"dependant_pct": 50},
        "GMM": {"employee_pct": 100, "dependant_pct": 100},
    }
    config["rules"] = [{"product_code": "GMM", "requires_product_code": "GHS"}]
    assert (
        broker.put(f"/api/v1/enrollment-windows/{wid}/form-config", json=config).status_code == 200
    )
    return wid


def _token(broker: TestClient, wid: str) -> str:
    assert broker.post(f"/api/v1/enrollment-windows/{wid}/qr-entry").status_code == 200
    url = broker.post(f"/api/v1/enrollment-windows/{wid}/qr-entry/link").json()["url"]
    assert "/enrol#" in url
    return url.split("#", 1)[1]


def _hdr(token: str) -> dict[str, str]:
    return {"X-Inspro-Link": token}


def _clauses(public: TestClient, token: str) -> list[str]:
    return [
        c["id"] for c in public.get("/api/v1/public/enrol", headers=_hdr(token)).json()["clauses"]
    ]


def _payload(token_clauses: list[str], **over: Any) -> dict[str, Any]:
    body: dict[str, Any] = {
        "request_id": str(uuid4()),
        "particulars": {
            "name": "Tan Ah Kow",
            "id_no": EMP_NRIC,
            "staff_id": "Q-1",
            "gender": "M",
            "contact_no": "91234567",
            "email": "ahkow@example.test",
        },
        "dependants": [
            {
                "key": "d1",
                "action": "add",
                "relationship": "spouse",
                "name": "Lim Mei Ling",
                "gender": "F",
                "id_no": SPOUSE_NRIC,
                "dob": "1985-05-05",
            }
        ],
        "products": [
            {
                "product_code": "GHS",
                "category_id": "qr-cat-senior",
                "action": "enrol",
                "tier_key": "qr-cat-senior::P1",
                "covered": ["d1"],
            },
            {
                "product_code": "GMM",
                "category_id": "qr-cat-gmm",
                "action": "enrol",
                "tier_key": "qr-cat-gmm::M1",
                "covered": ["d1"],
            },
        ],
        "accepted_clause_ids": token_clauses,
        "signature_name": "tan ah  kow",
    }
    body.update(over)
    return body


def _submit(public: TestClient, token: str, **over: Any):
    return public.post(
        "/api/v1/public/enrol/submissions",
        headers=_hdr(token),
        json=_payload(_clauses(public, token), **over),
    )


# ── Shared employee QR ───────────────────────────────────────────────────────


def test_public_form_lists_every_cohort_with_the_employee_share(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    res = TestClient(app).get("/api/v1/public/enrol", headers=_hdr(token))
    assert res.status_code == 200
    assert res.headers["cache-control"] == "no-store"
    form = res.json()
    assert form["status"] == "open"
    assert form["company_name"] == "QR Test Company Pte Ltd"
    ghs = next(p for p in form["products"] if p["product_code"] == "GHS")
    assert {c["label"] for c in ghs["cohorts"]} == {"Senior management", "All other staff"}
    senior = next(
        t for c in form["contributions"] for t in c["tiers"] if t["tier_key"] == "qr-cat-senior::P1"
    )
    assert senior["premium"] == {"EO": 400.0, "ES": 1000.0, "EC": 900.0, "EF": 1500.0}
    assert senior["family"]["spouse"] == 300.0  # 50% of (1000 - 400)
    assert {"qr_identity", "qr_personal_data"} <= {c["id"] for c in form["clauses"]}


def test_unknown_or_missing_link_is_404(broker: TestClient) -> None:
    public = TestClient(app)
    assert public.get("/api/v1/public/enrol").status_code == 404
    assert public.get("/api/v1/public/enrol", headers=_hdr("x" * 43)).status_code == 404


def test_paused_and_rotated_links(broker: TestClient) -> None:
    wid = _window(broker)
    token = _token(broker, wid)
    public = TestClient(app)
    broker.patch(f"/api/v1/enrollment-windows/{wid}/qr-entry", json={"enabled": False})
    paused = public.get("/api/v1/public/enrol", headers=_hdr(token)).json()
    assert paused["status"] == "paused" and paused["products"] == []
    broker.patch(f"/api/v1/enrollment-windows/{wid}/qr-entry", json={"enabled": True})
    broker.post(f"/api/v1/enrollment-windows/{wid}/qr-entry/rotate")
    assert public.get("/api/v1/public/enrol", headers=_hdr(token)).status_code == 404
    fresh = broker.post(f"/api/v1/enrollment-windows/{wid}/qr-entry/link").json()["url"]
    again = broker.post(f"/api/v1/enrollment-windows/{wid}/qr-entry/link").json()["url"]
    assert fresh == again  # re-download gives the same QR
    assert public.get("/api/v1/public/enrol", headers=_hdr(fresh.split("#")[1])).status_code == 200


def test_viewer_cannot_turn_on_or_reveal_the_qr(broker: TestClient) -> None:
    wid = _window(broker)
    app.dependency_overrides[get_current_user] = lambda: _broker("broker_viewer")
    viewer = TestClient(app)
    assert viewer.get(f"/api/v1/enrollment-windows/{wid}/qr-entry").status_code == 200
    assert viewer.post(f"/api/v1/enrollment-windows/{wid}/qr-entry").status_code == 403
    assert viewer.post(f"/api/v1/enrollment-windows/{wid}/qr-entry/link").status_code == 403


def test_submission_is_pending_and_changes_no_roster_data(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    with SessionLocal() as s:
        before = (
            s.scalar(select(func.count(Employee.id))),
            s.scalar(select(func.count(Dependant.id))),
        )
    res = _submit(TestClient(app), token)
    assert res.status_code == 201, res.text
    receipt = res.json()
    assert receipt["reference_no"] == "QR-2029-00001"
    with SessionLocal() as s:
        assert (
            s.scalar(select(func.count(Employee.id))),
            s.scalar(select(func.count(Dependant.id))),
        ) == before
        intake = s.scalars(select(PublicIntake)).one()
        assert intake.status == "pending"
        ghs, gmm = intake.snapshot["selections"]
        assert ghs["composition"] == "ES" and ghs["you_pay"] == 300.0
        assert gmm["you_pay"] == 270.0 and gmm["full_premium"] == 270.0
        doc = s.get(StoredDocument, intake.document_id)
        assert doc is not None and doc.entity_type == "qr_intake"
        audit = s.scalars(select(AuditLog).where(AuditLog.action == "qr_intake.submitted")).one()
        assert audit.actor_type == "public" and audit.user_id is None


def test_retry_is_idempotent_and_a_changed_body_is_refused(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    public = TestClient(app)
    body = _payload(_clauses(public, token))
    first = public.post("/api/v1/public/enrol/submissions", headers=_hdr(token), json=body)
    again = public.post("/api/v1/public/enrol/submissions", headers=_hdr(token), json=body)
    assert first.json()["reference_no"] == again.json()["reference_no"]
    body["particulars"]["contact_no"] = "98765432"
    changed = public.post("/api/v1/public/enrol/submissions", headers=_hdr(token), json=body)
    assert changed.status_code == 409
    with SessionLocal() as s:
        assert s.scalar(select(func.count(PublicIntake.id))) == 1


def test_server_checks_every_rule(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    public = TestClient(app)
    clauses = _clauses(public, token)
    bad = _payload(clauses[:-1], signature_name="Someone Else")
    bad["particulars"]["id_no"] = "S1234567A"
    bad["dependants"][0]["dob"] = "1940-01-01"  # beyond the spouse age window
    bad["products"][0]["covered"] = []  # GMM covers the spouse, GHS does not
    res = public.post("/api/v1/public/enrol/submissions", headers=_hdr(token), json=bad)
    assert res.status_code == 422
    problems = " ".join(res.json()["detail"]["problems"])
    assert "valid NRIC" in problems
    assert "Type your full name" in problems
    assert "Tick every declaration" in problems
    assert "must also be covered under" in problems
    assert "outside the age limit" in problems
    honeypot = public.post(
        "/api/v1/public/enrol/submissions",
        headers=_hdr(token),
        json=_payload(clauses, website="http://spam.test"),
    )
    assert honeypot.status_code == 400


def test_own_copy_downloads_once(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    public = TestClient(app)
    receipt = _submit(public, token).json()
    body = {"reference_no": receipt["reference_no"], "copy_key": receipt["copy_key"]}
    wrong = public.post(
        "/api/v1/public/enrol/copy", headers=_hdr(token), json={**body, "copy_key": "x" * 32}
    )
    assert wrong.status_code == 404
    first = public.post("/api/v1/public/enrol/copy", headers=_hdr(token), json=body)
    assert first.status_code == 200 and first.content.startswith(b"%PDF")
    assert (
        public.post("/api/v1/public/enrol/copy", headers=_hdr(token), json=body).status_code == 404
    )


def test_closed_period_refuses_submissions(broker: TestClient) -> None:
    wid = _window(broker)
    token = _token(broker, wid)
    with SessionLocal() as s:
        s.get(EnrollmentWindow, wid).closes_at = datetime.now(UTC) - timedelta(days=1)
        s.commit()
    public = TestClient(app)
    assert public.get("/api/v1/public/enrol", headers=_hdr(token)).json()["status"] == "closed"
    res = public.post(
        "/api/v1/public/enrol/submissions", headers=_hdr(token), json=_payload(["qr_identity"])
    )
    assert res.status_code == 409


# ── Broker review ────────────────────────────────────────────────────────────


def test_link_files_the_form_like_a_paper_form(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    _submit(TestClient(app), token)
    listed = broker.get(f"/api/v1/policy-years/{PY_ID}/qr-intakes").json()
    assert listed["counts"] == {"pending": 1}
    intake_id = listed["items"][0]["id"]
    detail = broker.get(f"/api/v1/qr-intakes/{intake_id}").json()
    best = detail["candidates"][0]
    assert best["employee_id"] == EMP1
    assert set(best["match"]) == {"Staff ID", "NRIC / FIN", "Name"}
    linked = broker.post(f"/api/v1/qr-intakes/{intake_id}/link", json={"employee_id": EMP1})
    assert linked.status_code == 200, linked.text
    assert linked.json()["source"] == "qr" and linked.json()["reference_no"] == "QR-2029-00001"
    again = broker.post(f"/api/v1/qr-intakes/{intake_id}/link", json={"employee_id": EMP1})
    assert again.status_code == 409
    register = broker.get(f"/api/v1/policy-years/{PY_ID}/enrollment-forms?source=qr").json()
    assert register["total"] == 1
    pdf = broker.get(f"/api/v1/enrollment-forms/{linked.json()['id']}/pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def test_register_workbook_has_qr_sheets(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    _submit(TestClient(app), token)
    res = broker.get(f"/api/v1/policy-years/{PY_ID}/enrollment-forms/export.xlsx")
    wb = load_workbook(io.BytesIO(res.content))
    assert {"QR submissions", "Dependants enrolment"} <= set(wb.sheetnames)
    rows = list(wb["Dependants enrolment"].iter_rows(values_only=True))
    actions = [(r[0], r[5], r[16]) for r in rows[1:]]
    assert ("", "Employee", "ES") in actions or (None, "Employee", "ES") in actions
    assert ("A", "Spouse", None) in actions or ("A", "Spouse", "") in actions
    masked_id = rows[1][8]
    assert masked_id and "*" in masked_id


def test_reject_and_restore(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    _submit(TestClient(app), token)
    intake_id = broker.get(f"/api/v1/policy-years/{PY_ID}/qr-intakes").json()["items"][0]["id"]
    assert (
        broker.post(f"/api/v1/qr-intakes/{intake_id}/reject", json={"reason": "other"}).status_code
        == 422
    )
    assert (
        broker.post(
            f"/api/v1/qr-intakes/{intake_id}/reject", json={"reason": "duplicate"}
        ).status_code
        == 204
    )
    assert (
        broker.post(f"/api/v1/qr-intakes/{intake_id}/link", json={"employee_id": EMP1}).status_code
        == 409
    )
    assert broker.post(f"/api/v1/qr-intakes/{intake_id}/restore").status_code == 204


def test_other_companies_submissions_are_invisible(broker: TestClient) -> None:
    with SessionLocal() as s:
        foreign = PublicIntake(
            client_id=OTHER_CLIENT,
            policy_year_id=PY_ID,
            entry_id="x",
            reference_no="QR-1",
            snapshot={},
            content_sha256="0" * 64,
            request_id="r" * 8,
            request_hash="h",
        )
        s.add(foreign)
        s.commit()
        fid = foreign.id
    assert broker.get(f"/api/v1/qr-intakes/{fid}").status_code == 404
    assert (
        broker.post(f"/api/v1/qr-intakes/{fid}/reject", json={"reason": "duplicate"}).status_code
        == 404
    )


def test_close_waits_for_pending_qr_forms(broker: TestClient) -> None:
    wid = _window(broker)
    _submit(TestClient(app), _token(broker, wid))
    preview = broker.get(f"/api/v1/enrollment-windows/{wid}/close-preview").json()
    assert preview["pending_qr_intakes"] == 1
    refused = broker.post(f"/api/v1/enrollment-windows/{wid}/close")
    assert refused.status_code == 409
    assert refused.json()["detail"]["code"] == "pending_qr_intakes"
    ok = broker.post(
        f"/api/v1/enrollment-windows/{wid}/close", json={"close_with_pending_qr": True}
    )
    assert ok.status_code == 200, ok.text


# ── HR download: batch, link + PIN, ten minutes ──────────────────────────────


def _batch_with_link(broker: TestClient) -> tuple[str, str, str]:
    token = _token(broker, _window(broker))
    _submit(TestClient(app), token)
    batch = broker.post(
        f"/api/v1/policy-years/{PY_ID}/hr-batches", json={"include_pending_qr": True}
    )
    assert batch.status_code == 201, batch.text
    assert batch.json()["file_count"] == 1
    issued = broker.post(
        f"/api/v1/hr-batches/{batch.json()['id']}/grants", json={"recipient_note": "HR team"}
    ).json()
    assert issued["pin"].isdigit() and len(issued["pin"]) == 6
    return batch.json()["id"], issued["url"].split("#", 1)[1], issued["pin"]


def _wrong(pin: str) -> str:
    return f"{(int(pin) + 1) % 1_000_000:06d}"


def test_pin_then_start_gives_one_browser_ten_minutes(broker: TestClient, monkeypatch) -> None:
    _batch_id, token, pin = _batch_with_link(broker)
    hr = TestClient(app)
    path = "/api/v1/public/hr-download"
    assert hr.get(path, headers=_hdr(token)).json()["state"] == "ready"
    assert hr.post(f"{path}/start", headers=_hdr(token)).status_code == 401
    wrong = hr.post(f"{path}/pin", headers=_hdr(token), json={"pin": _wrong(pin)})
    assert wrong.status_code == 401 and wrong.json()["detail"]["attempts_left"] == 4
    assert (
        hr.post(f"{path}/pin", headers=_hdr(token), json={"pin": pin}).json()["state"] == "verified"
    )
    assert hr.get(f"{path}/file").status_code == 409
    started = hr.post(f"{path}/start", headers=_hdr(token)).json()
    assert started["state"] == "active"
    deadline = datetime.fromisoformat(started["expires_at"])
    retry = hr.post(f"{path}/start", headers=_hdr(token)).json()
    assert datetime.fromisoformat(retry["expires_at"]) == deadline  # retry keeps the deadline
    zip_res = hr.get(f"{path}/file")
    assert zip_res.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(zip_res.content)).namelist()
    assert any(n.startswith("QR-2029-00001 PENDING VERIFICATION") for n in names)
    assert "Forms in this download.xlsx" in names
    other = TestClient(app)
    other.post(f"{path}/pin", headers=_hdr(token), json={"pin": pin})
    assert other.get(path, headers=_hdr(token)).json()["state"] == "in_use"
    assert other.get(f"{path}/file").status_code == 404
    aware_deadline = deadline if deadline.tzinfo else deadline.replace(tzinfo=UTC)
    monkeypatch.setattr(batch_svc, "now_utc", lambda: aware_deadline + timedelta(seconds=1))
    assert hr.get(f"{path}/file").status_code == 409
    assert hr.get(path, headers=_hdr(token)).json()["state"] == "expired"


def test_five_wrong_pins_lock_and_reissue_revokes(broker: TestClient) -> None:
    batch_id, token, pin = _batch_with_link(broker)
    hr = TestClient(app)
    path = "/api/v1/public/hr-download"
    for _ in range(4):
        assert (
            hr.post(f"{path}/pin", headers=_hdr(token), json={"pin": _wrong(pin)}).status_code
            == 401
        )
    locked = hr.post(f"{path}/pin", headers=_hdr(token), json={"pin": _wrong(pin)})
    assert locked.status_code == 409 and locked.json()["detail"]["state"] == "locked"
    assert hr.post(f"{path}/pin", headers=_hdr(token), json={"pin": pin}).status_code == 409
    fresh = broker.post(f"/api/v1/hr-batches/{batch_id}/grants", json={}).json()
    assert hr.get(path, headers=_hdr(token)).json()["state"] == "revoked"
    new_token = fresh["url"].split("#", 1)[1]
    assert hr.get(path, headers=_hdr(new_token)).json()["state"] == "ready"


def test_activation_race_has_one_winner(broker: TestClient) -> None:
    _batch_id, token, _pin = _batch_with_link(broker)
    with SessionLocal() as s:
        grant, _batch = batch_svc.resolve_grant(s, token, None)
        assert batch_svc.activate(s, grant, "first-browser") is True
        assert batch_svc.activate(s, grant, "second-browser") is False
        assert batch_svc.activate(s, grant, "first-browser") is True
        s.commit()


def test_batch_with_a_missing_pdf_is_never_published(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    _submit(TestClient(app), token)
    with SessionLocal() as s:
        s.query(StoredDocument).delete()
        s.commit()
    res = broker.post(f"/api/v1/policy-years/{PY_ID}/hr-batches", json={"include_pending_qr": True})
    assert res.status_code == 409
    assert res.json()["detail"]["code"] == "batch_files_missing"
    with SessionLocal() as s:
        assert s.scalar(select(func.count(HrDownloadBatch.id))) == 0


def test_only_bulk_export_roles_prepare_and_issue(broker: TestClient) -> None:
    app.dependency_overrides[get_current_user] = lambda: replace(_broker(), role="broker_viewer")
    viewer = TestClient(app)
    assert (
        viewer.post(
            f"/api/v1/policy-years/{PY_ID}/hr-batches", json={"include_pending_qr": True}
        ).status_code
        == 403
    )


# ── Hardening found by the security review ───────────────────────────────────


def test_control_characters_are_refused(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    public = TestClient(app)
    body = _payload(_clauses(public, token))
    body["particulars"]["name"] = "Tan\x00Ah Kow"
    res = public.post("/api/v1/public/enrol/submissions", headers=_hdr(token), json=body)
    assert res.status_code == 422
    body = _payload(_clauses(public, token), accepted_clause_ids=["x" * 500])
    res = public.post("/api/v1/public/enrol/submissions", headers=_hdr(token), json=body)
    assert res.status_code == 422


def test_one_person_cannot_flood_the_queue(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    public = TestClient(app)
    for _ in range(3):
        assert _submit(public, token).status_code == 201
    fourth = _submit(public, token)
    assert fourth.status_code == 429
    assert fourth.json()["detail"]["code"] == "already_submitted"


def test_a_link_only_works_on_its_own_firms_host(broker: TestClient) -> None:
    from app.models import BrokerFirm, TenantDomain

    token = _token(broker, _window(broker))
    with SessionLocal() as s:
        s.add(
            BrokerFirm(id="qr-other-firm", name="Other Firm", slug="qr-other-firm", status="active")
        )
        s.flush()
        s.add(
            TenantDomain(
                broker_firm_id="qr-other-firm",
                hostname="other.qr.test",
                surface="all",
                is_primary=True,
                status="active",
            )
        )
        s.commit()
    try:
        other_host = TestClient(app, base_url="http://other.qr.test")
        assert other_host.get("/api/v1/public/enrol", headers=_hdr(token)).status_code == 404
        assert TestClient(app).get("/api/v1/public/enrol", headers=_hdr(token)).status_code == 200
    finally:
        with SessionLocal() as s:
            s.query(TenantDomain).filter(TenantDomain.hostname == "other.qr.test").delete()
            s.query(BrokerFirm).filter(BrokerFirm.id == "qr-other-firm").delete()
            s.commit()


def test_linking_over_an_existing_form_needs_confirmation(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    public = TestClient(app)
    _submit(public, token)
    _submit(public, token)
    first, second = sorted(
        broker.get(f"/api/v1/policy-years/{PY_ID}/qr-intakes").json()["items"],
        key=lambda i: i["reference_no"],
    )
    assert (
        broker.post(
            f"/api/v1/qr-intakes/{first['id']}/link", json={"employee_id": EMP1}
        ).status_code
        == 200
    )
    refused = broker.post(f"/api/v1/qr-intakes/{second['id']}/link", json={"employee_id": EMP1})
    assert refused.status_code == 409
    assert refused.json()["detail"]["code"] == "existing_form"
    ok = broker.post(
        f"/api/v1/qr-intakes/{second['id']}/link",
        json={"employee_id": EMP1, "confirm_supersede": True},
    )
    assert ok.status_code == 200


def test_downloads_per_link_are_capped(broker: TestClient) -> None:
    _batch_id, token, pin = _batch_with_link(broker)
    hr = TestClient(app)
    path = "/api/v1/public/hr-download"
    hr.post(f"{path}/pin", headers=_hdr(token), json={"pin": pin})
    hr.post(f"{path}/start", headers=_hdr(token))
    for _ in range(batch_svc.MAX_DOWNLOADS):
        assert hr.get(f"{path}/file").status_code == 200
    capped = hr.get(f"{path}/file")
    assert capped.status_code == 409
    assert capped.json()["detail"]["code"] == "download_limit"
    assert hr.get(path, headers=_hdr(token)).json()["downloads_left"] == 0


def test_public_form_carries_no_raw_financials(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    form = TestClient(app).get("/api/v1/public/enrol", headers=_hdr(token)).json()
    tiers = [t for p in form["products"] for c in p["cohorts"] for t in c["tier_set"]["tiers"]]
    assert tiers and all(t["financials"] is None for t in tiers)
    assert form["contributions"]  # premiums still shown, from the computed table


def test_preview_rejects_unknown_status(broker: TestClient) -> None:
    res = broker.get(f"/api/v1/policy-years/{PY_ID}/hr-batches/preview?statuses=bogus")
    assert res.status_code == 422


# ── Regressions found by code review ─────────────────────────────────────────


def test_uploaded_document_opens_from_the_qr_form(broker: TestClient) -> None:
    wid = _window(broker)
    uploaded = broker.post(
        f"/api/v1/enrollment-windows/{wid}/form-config/documents",
        files={"file": ("guide.pdf", b"%PDF-1.4 product guide", "application/pdf")},
    )
    assert uploaded.status_code == 200, uploaded.text
    config = f"/api/v1/enrollment-windows/{wid}/form-config"
    settings = broker.get(config).json()["settings"]
    settings["documents"] = [
        {"id": "guide", "label": "Product guide", "document_id": uploaded.json()["document_id"]}
    ]
    assert broker.put(config, json=settings).status_code == 200
    token = _token(broker, wid)
    public = TestClient(app)
    res = public.get("/api/v1/public/enrol/documents/guide", headers=_hdr(token))
    assert res.status_code == 200 and res.content == b"%PDF-1.4 product guide"
    stored_id = uploaded.json()["document_id"]
    assert (
        public.get(f"/api/v1/public/enrol/documents/{stored_id}", headers=_hdr(token)).status_code
        == 404
    )


def test_rejected_submissions_stay_out_of_hr_exports(broker: TestClient) -> None:
    from app.services.enrollment_forms.register import RegisterFilter, build_workbook

    token = _token(broker, _window(broker))
    _submit(TestClient(app), token)
    intake_id = broker.get(f"/api/v1/policy-years/{PY_ID}/qr-intakes").json()["items"][0]["id"]
    broker.post(f"/api/v1/qr-intakes/{intake_id}/reject", json={"reason": "wrong_company"})
    with SessionLocal() as s:
        hr = build_workbook(s, RegisterFilter(client_id=CLIENT_ID), include_rejected_qr=False)
        mine = build_workbook(s, RegisterFilter(client_id=CLIENT_ID))
    assert hr.qr_submissions == 0
    assert "QR submissions" not in load_workbook(io.BytesIO(hr.content)).sheetnames
    assert mine.qr_submissions == 1


def test_missing_blob_fails_the_batch_with_a_clear_error(broker: TestClient) -> None:
    from app.core.storage import get_storage

    token = _token(broker, _window(broker))
    _submit(TestClient(app), token)
    with SessionLocal() as s:
        doc = s.scalars(select(StoredDocument)).one()
        get_storage().delete(doc.storage_path)
    res = broker.post(f"/api/v1/policy-years/{PY_ID}/hr-batches", json={"include_pending_qr": True})
    assert res.status_code == 409
    assert res.json()["detail"]["code"] == "batch_files_missing"


def test_life_event_dates_do_not_move_age_checks_in_an_annual_period(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    public = TestClient(app)
    body = _payload(_clauses(public, token), life_event={"kind": "other", "date": "2000-01-01"})
    body["dependants"][0]["dob"] = "1940-01-01"
    res = public.post("/api/v1/public/enrol/submissions", headers=_hdr(token), json=body)
    assert res.status_code == 422
    assert "outside the age limit" in " ".join(res.json()["detail"]["problems"])


def test_an_older_qr_form_never_supersedes_a_newer_one(broker: TestClient) -> None:
    token = _token(broker, _window(broker))
    public = TestClient(app)
    _submit(public, token)
    _submit(public, token)
    older, newer = sorted(
        broker.get(f"/api/v1/policy-years/{PY_ID}/qr-intakes").json()["items"],
        key=lambda i: i["reference_no"],
    )
    with SessionLocal() as s:
        row = s.get(PublicIntake, older["id"])
        row.created_at = datetime.now(UTC) - timedelta(hours=2)
        s.commit()
    assert (
        broker.post(
            f"/api/v1/qr-intakes/{newer['id']}/link", json={"employee_id": EMP1}
        ).status_code
        == 200
    )
    refused = broker.post(
        f"/api/v1/qr-intakes/{older['id']}/link",
        json={"employee_id": EMP1, "confirm_supersede": True},
    )
    assert refused.status_code == 409
    assert refused.json()["detail"]["code"] == "newer_form"
